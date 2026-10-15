"""Tests del lector de ALUMNOS (`importaciones/parsing/alumnos.py`).

`SimpleTestCase` con workbooks armados en memoria: el lector es Django-free y
no necesita base. El `hoy` es fijo, lejano a la fecha real, para que ningún
test dependa del día en que corre.
"""

import io
import inspect
from datetime import date, datetime

import openpyxl
from django.test import SimpleTestCase

from alumnos.ciclo import rango_inicio_ciclo
from importaciones import parsing
from importaciones.parsing import (
    ALIAS_ALUMNOS,
    MARCA_FILA_EJEMPLO,
    MAX_FILAS_ALUMNOS,
    ArchivoCSVIlegible,
    ColumnaRequeridaFaltante,
    DemasiadasFilas,
    leer_hoja_alumnos,
    libro_desde_csv,
    parsear_archivo_alumnos,
)
from importaciones.parsing import alumnos as lector

HOY = date(2026, 6, 15)


def hoja(*filas, merges=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    for fila in filas:
        ws.append(list(fila))
    for rango in merges:
        ws.merge_cells(rango)
    return ws


def leer(*filas, **kw):
    return leer_hoja_alumnos(hoja(*filas, **kw), hoy=HOY)


def una_fila(encabezados, valores):
    resultado = leer(encabezados, valores)
    assert len(resultado.filas) == 1, resultado.filas_invalidas
    return resultado.filas[0]


class FechasTests(SimpleTestCase):
    def test_celda_datetime(self):
        self.assertEqual(
            lector.leer_fecha(datetime(1990, 5, 3, 14, 30), hoy=HOY),
            (date(1990, 5, 3), None),
        )

    def test_numero_de_serie_de_excel(self):
        # 33000 = 1990-05-07 en el calendario de Excel.
        fecha, aviso = lector.leer_fecha(33000, hoy=HOY)
        self.assertEqual(fecha, date(1990, 5, 7))
        self.assertIsNone(aviso)

    def test_texto_dia_primero_con_cualquier_separador(self):
        for texto in ("03/05/1990", "3-5-1990", "3.5.1990"):
            with self.subTest(texto=texto):
                self.assertEqual(
                    lector.leer_fecha(texto, hoy=HOY), (date(1990, 5, 3), None)
                )

    def test_anio_de_dos_digitos_usa_el_corte_de_hoy(self):
        # hoy.year % 100 == 26: 90 > 26 -> 1990; 05 <= 26 -> 2005.
        self.assertEqual(lector.leer_fecha("3-5-90", hoy=HOY)[0], date(1990, 5, 3))
        self.assertEqual(lector.leer_fecha("3-5-05", hoy=HOY)[0], date(2005, 5, 3))

    def test_iso_con_hora(self):
        self.assertEqual(
            lector.leer_fecha("1990-05-03 14:30:00", hoy=HOY), (date(1990, 5, 3), None)
        )
        self.assertEqual(
            lector.leer_fecha("1990-05-03T14:30:00", hoy=HOY), (date(1990, 5, 3), None)
        )

    def test_fecha_inexistente_avisa(self):
        fecha, aviso = lector.leer_fecha("31/02/1990", hoy=HOY)
        self.assertIsNone(fecha)
        self.assertTrue(aviso)

    def test_anio_suelto_avisa(self):
        for valor in (1990, "1990"):
            with self.subTest(valor=valor):
                fecha, aviso = lector.leer_fecha(valor, hoy=HOY)
                self.assertIsNone(fecha)
                self.assertIn("año", aviso)

    def test_vacio_no_avisa(self):
        self.assertEqual(lector.leer_fecha(None, hoy=HOY), (None, None))
        self.assertEqual(lector.leer_fecha("  ", hoy=HOY), (None, None))

    def test_texto_cualquiera_avisa(self):
        self.assertIsNone(lector.leer_fecha("mañana", hoy=HOY)[0])

    def test_nacimiento_futuro_o_antes_de_1900_se_rechaza(self):
        fecha, aviso = lector.leer_fecha_nacimiento(date(2026, 6, 16), hoy=HOY)
        self.assertIsNone(fecha)
        self.assertTrue(aviso)
        self.assertIsNone(lector.leer_fecha_nacimiento(date(1899, 12, 31), hoy=HOY)[0])
        # Hoy mismo es válido.
        self.assertEqual(lector.leer_fecha_nacimiento(HOY, hoy=HOY)[0], HOY)

    def test_fila_con_nacimiento_futuro_lo_deja_vacio_y_lo_anota(self):
        fila = una_fila(("Nombre", "Apellido", "Nacimiento"), ("Ana", "Gómez", "01/01/2030"))
        self.assertIsNone(fila.datos["fecha_nacimiento"])
        self.assertEqual(fila.no_cargados, ["Fecha de nacimiento «01/01/2030»"])
        self.assertTrue(fila.avisos)

    def test_fila_con_anio_suelto_lo_anota_con_el_original(self):
        fila = una_fila(("Nombre", "Apellido", "Nacimiento"), ("Ana", "Gómez", 1990))
        self.assertIsNone(fila.datos["fecha_nacimiento"])
        self.assertEqual(fila.no_cargados, ["Fecha de nacimiento «1990»"])


class InicioDelCicloTests(SimpleTestCase):
    def test_rango_es_hoy_menos_27_a_hoy_mas_60(self):
        desde, hasta = rango_inicio_ciclo(HOY)
        self.assertEqual((HOY - desde).days, 27)
        self.assertEqual((hasta - HOY).days, 60)

    def test_los_bordes_pasan(self):
        desde, hasta = rango_inicio_ciclo(HOY)
        self.assertEqual(lector.leer_inicio_ciclo(desde, hoy=HOY), (desde, None))
        self.assertEqual(lector.leer_inicio_ciclo(hasta, hoy=HOY), (hasta, None))

    def test_un_dia_afuera_queda_vacio_con_aviso(self):
        desde, hasta = rango_inicio_ciclo(HOY)
        for fecha in (date.fromordinal(desde.toordinal() - 1),
                      date.fromordinal(hasta.toordinal() + 1)):
            with self.subTest(fecha=fecha):
                resultado, aviso = lector.leer_inicio_ciclo(fecha, hoy=HOY)
                self.assertIsNone(resultado)
                self.assertTrue(aviso)

    def test_fila_fuera_de_rango_lo_anota(self):
        fila = una_fila(
            ("Nombre", "Apellido", "Inicio del ciclo"), ("Ana", "Gómez", "01/01/2020")
        )
        self.assertIsNone(fila.datos["fecha_inicio_ciclo"])
        self.assertEqual(fila.no_cargados, ["Inicio del ciclo «01/01/2020»"])

    def test_fila_dentro_del_rango(self):
        fila = una_fila(
            ("Nombre", "Apellido", "Inicio del ciclo"), ("Ana", "Gómez", "20/06/2026")
        )
        self.assertEqual(fila.datos["fecha_inicio_ciclo"], date(2026, 6, 20))
        self.assertEqual(fila.no_cargados, [])

    def test_fecha_de_alta_no_es_inicio_de_ciclo(self):
        resultado = leer(("Nombre", "Apellido", "Fecha de alta"), ("Ana", "Gómez", "01/01/2020"))
        self.assertNotIn("inicio_ciclo", resultado.columnas_detectadas)
        self.assertIn("Fecha de alta", resultado.columnas_ignoradas)
        self.assertIsNone(resultado.filas[0].datos["fecha_inicio_ciclo"])


class EstadoTests(SimpleTestCase):
    def test_activos(self):
        for valor in ("Activo", "activa", "Alta", "vigente", "Sí", "si"):
            with self.subTest(valor=valor):
                self.assertEqual(lector.leer_estado(valor), ("activo", None))

    def test_bajas(self):
        for valor in ("Inactivo", "inactiva", "Baja", "de baja", "dado de baja",
                      "Dada de baja", "No"):
            with self.subTest(valor=valor):
                self.assertEqual(lector.leer_estado(valor), ("inactivo", None))

    def test_vacio_es_activo_sin_aviso(self):
        self.assertEqual(lector.leer_estado(None), ("activo", None))
        self.assertEqual(lector.leer_estado("  "), ("activo", None))

    def test_ilegible_es_activo_con_aviso_y_se_anota(self):
        self.assertEqual(lector.leer_estado("quién sabe")[0], "activo")
        self.assertTrue(lector.leer_estado("quién sabe")[1])
        fila = una_fila(("Nombre", "Apellido", "Estado"), ("Ana", "Gómez", "quién sabe"))
        self.assertEqual(fila.datos["estado"], "activo")
        self.assertEqual(fila.no_cargados, ["Estado «quién sabe»"])


class BooleanosTests(SimpleTestCase):
    def test_tabla_completa(self):
        for valor in ("sí", "si", "s", "x", "X", "1", 1, "true", "verdadero", "yes", "y", "✓", True):
            with self.subTest(verdadero=valor):
                self.assertEqual(lector.leer_booleano(valor), (True, ""))
        for valor in ("no", "n", "0", 0, "false", "falso", "", None, "-", False):
            with self.subTest(falso=valor):
                self.assertEqual(lector.leer_booleano(valor), (False, ""))

    def test_si_con_resto(self):
        self.assertEqual(lector.leer_booleano("Sí, asma"), (True, "asma"))
        self.assertEqual(lector.leer_booleano("si - diabetes"), (True, "diabetes"))

    def test_una_palabra_que_empieza_con_si_no_es_un_si(self):
        self.assertEqual(lector.leer_booleano("silla de ruedas"), (None, ""))

    def test_texto_no_entendido(self):
        self.assertEqual(lector.leer_booleano("tal vez"), (None, ""))

    def test_si_con_resto_pasa_el_resto_al_detalle(self):
        fila = una_fila(
            ("Nombre", "Apellido", "Enfermedad crónica", "Detalle enfermedad crónica"),
            ("Ana", "Gómez", "Sí, asma", None),
        )
        self.assertTrue(fila.datos["tiene_enfermedad_cronica"])
        self.assertEqual(fila.datos["enfermedad_cronica_detalle"], "asma")
        self.assertTrue(fila.avisos)

    def test_si_con_resto_no_pisa_un_detalle_ya_cargado(self):
        fila = una_fila(
            ("Nombre", "Apellido", "Enfermedad crónica", "Detalle enfermedad crónica"),
            ("Ana", "Gómez", "Sí, asma", "diabetes"),
        )
        self.assertEqual(fila.datos["enfermedad_cronica_detalle"], "diabetes")
        self.assertEqual(fila.no_cargados, ["Enfermedad crónica «Sí, asma»"])

    def test_detalle_sin_tiene_lo_infiere(self):
        fila = una_fila(
            ("Nombre", "Apellido", "Detalle discapacidad"), ("Ana", "Gómez", "Silla de ruedas")
        )
        self.assertTrue(fila.datos["tiene_discapacidad"])
        self.assertTrue(fila.avisos)

    def test_detalle_con_tiene_vacio_lo_infiere(self):
        fila = una_fila(
            ("Nombre", "Apellido", "Tiene discapacidad", "Detalle discapacidad"),
            ("Ana", "Gómez", None, "Sordera"),
        )
        self.assertTrue(fila.datos["tiene_discapacidad"])

    def test_detalle_ninguna_no_infiere(self):
        for detalle in ("ninguna", "Ninguno", "-", "n/a", "NA", "no"):
            with self.subTest(detalle=detalle):
                fila = una_fila(
                    ("Nombre", "Apellido", "Detalle discapacidad"), ("Ana", "Gómez", detalle)
                )
                self.assertFalse(fila.datos["tiene_discapacidad"])
                self.assertEqual(fila.avisos, [])

    def test_tiene_explicito_en_no_no_se_pisa(self):
        fila = una_fila(
            ("Nombre", "Apellido", "Tiene discapacidad", "Detalle discapacidad"),
            ("Ana", "Gómez", "No", "Sordera"),
        )
        self.assertFalse(fila.datos["tiene_discapacidad"])

    def test_columna_unica_con_texto_libre(self):
        fila = una_fila(("Nombre", "Apellido", "Discapacidad"), ("Ana", "Gómez", "Sordera"))
        self.assertTrue(fila.datos["tiene_discapacidad"])
        self.assertEqual(fila.datos["discapacidad_detalle"], "Sordera")

    def test_frecuencia_con_valor_infiere_actividad_previa(self):
        fila = una_fila(
            ("Nombre", "Apellido", "Frecuencia"), ("Ana", "Gómez", "diaria")
        )
        self.assertTrue(fila.datos["actividad_fisica_previa"])
        self.assertEqual(fila.datos["frecuencia_actividad_previa"], "diaria")
        self.assertTrue(fila.avisos)

    def test_actividad_en_no_con_frecuencia_no_se_pisa(self):
        fila = una_fila(
            ("Nombre", "Apellido", "Actividad física previa", "Frecuencia"),
            ("Ana", "Gómez", "No", "diaria"),
        )
        self.assertFalse(fila.datos["actividad_fisica_previa"])

    def test_actividad_no_entendida_es_falso_con_aviso_y_se_anota(self):
        fila = una_fila(
            ("Nombre", "Apellido", "Actividad física previa"), ("Ana", "Gómez", "tal vez")
        )
        self.assertFalse(fila.datos["actividad_fisica_previa"])
        self.assertEqual(fila.no_cargados, ["Actividad física previa «tal vez»"])


class SexoTests(SimpleTestCase):
    def test_palabras(self):
        for valor, esperado in (
            ("Masculino", "masculino"), ("hombre", "masculino"), ("Varón", "masculino"),
            ("masc", "masculino"), ("Femenino", "femenino"), ("Mujer", "femenino"),
            ("fem", "femenino"), ("Prefiere no decir", "no_decir"),
            ("no decir", "no_decir"), ("ND", "no_decir"),
        ):
            with self.subTest(valor=valor):
                self.assertEqual(lector.leer_sexo(valor), (esperado, None))

    def test_f_y_h(self):
        self.assertEqual(lector.leer_sexo("F")[0], "femenino")
        self.assertEqual(lector.leer_sexo("H")[0], "masculino")

    def test_convencion_f_m(self):
        self.assertEqual(lector.convencion_sexo(["F", "M", "F"]), "masculino")

    def test_convencion_h_m(self):
        self.assertEqual(lector.convencion_sexo(["H", "M"]), "femenino")

    def test_m_ambigua(self):
        for columna in (["M", "M"], ["F", "H", "M"]):
            with self.subTest(columna=columna):
                self.assertIsNone(lector.convencion_sexo(columna))
        sexo, aviso = lector.leer_sexo("M")
        self.assertEqual(sexo, "")
        self.assertTrue(aviso)

    def test_la_convencion_se_decide_con_la_columna_entera(self):
        resultado = leer(
            ("Nombre", "Apellido", "Sexo"),
            ("Ana", "Gómez", "F"),
            ("Juan", "Pérez", "M"),
        )
        self.assertEqual(
            [f.datos["sexo"] for f in resultado.filas], ["femenino", "masculino"]
        )

    def test_m_sin_pistas_queda_vacia_y_se_anota(self):
        fila = una_fila(("Nombre", "Apellido", "Sexo"), ("Juan", "Pérez", "M"))
        self.assertEqual(fila.datos["sexo"], "")
        self.assertEqual(fila.no_cargados, ["Sexo «M»"])

    def test_x_queda_vacio_y_se_anota(self):
        fila = una_fila(("Nombre", "Apellido", "Sexo"), ("Juan", "Pérez", "X"))
        self.assertEqual(fila.datos["sexo"], "")
        self.assertEqual(fila.no_cargados, ["Sexo «X»"])


class FrecuenciaTests(SimpleTestCase):
    def test_etiquetas_y_valores(self):
        for valor, esperado in (
            ("Diaria", "diaria"), ("diario", "diaria"), ("todos los días", "diaria"),
            ("varias_por_semana", "varias_por_semana"),
            ("Varias veces por semana", "varias_por_semana"),
            ("Una vez por semana", "una_por_semana"), ("semanal", "una_por_semana"),
            ("una_por_semana", "una_por_semana"), ("Ocasional", "ocasional"),
            ("a veces", "ocasional"), ("esporádica", "ocasional"),
        ):
            with self.subTest(valor=valor):
                self.assertEqual(lector.leer_frecuencia(valor), (esperado, None))

    def test_veces_por_semana(self):
        for valor, esperado in (
            ("3 veces por semana", "varias_por_semana"),
            ("2/3 veces por semana", "varias_por_semana"),
            ("5x por semana", "varias_por_semana"),
            ("1 vez", "una_por_semana"), ("1 vez por semana", "una_por_semana"),
            ("7 veces por semana", "diaria"),
            (3, "varias_por_semana"), (1, "una_por_semana"), (7.0, "diaria"),
        ):
            with self.subTest(valor=valor):
                self.assertEqual(lector.leer_frecuencia(valor), (esperado, None))

    def test_ilegible_avisa_y_se_anota(self):
        self.assertEqual(lector.leer_frecuencia("de vez en cuando, che")[0], "")
        self.assertEqual(lector.leer_frecuencia(9)[0], "")
        fila = una_fila(("Nombre", "Apellido", "Frecuencia"), ("Ana", "Gómez", "ni idea"))
        self.assertEqual(fila.datos["frecuencia_actividad_previa"], "")
        self.assertEqual(fila.no_cargados, ["Frecuencia de actividad previa «ni idea»"])
        self.assertFalse(fila.datos["actividad_fisica_previa"])


class NombreTests(SimpleTestCase):
    def test_particion_con_coma(self):
        self.assertEqual(
            lector.partir_nombre_completo("Pérez, Juan Carlos"), ("Juan Carlos", "Pérez")
        )

    def test_particion_sin_coma_la_ultima_palabra_es_el_apellido(self):
        self.assertEqual(
            lector.partir_nombre_completo("María del Carmen López"),
            ("María del Carmen", "López"),
        )

    def test_particion_apellido_primero(self):
        self.assertEqual(
            lector.partir_nombre_completo("Pérez Juan Carlos", apellido_primero=True),
            ("Juan Carlos", "Pérez"),
        )

    def test_una_palabra_queda_sin_apellido(self):
        self.assertEqual(lector.partir_nombre_completo("Juan"), ("Juan", ""))

    def test_nombre_y_apellido_no_lo_toma_nombre(self):
        resultado = leer(("Nombre y apellido", "Email"), ("Juan Pérez", "j@x.com"))
        self.assertEqual(resultado.columnas_detectadas["nombre_completo"], "Nombre y apellido")
        self.assertNotIn("nombre", resultado.columnas_detectadas)
        fila = resultado.filas[0]
        self.assertEqual((fila.datos["nombre"], fila.datos["apellido"]), ("Juan", "Pérez"))
        self.assertIn("Separé nombre y apellido: revisalo", fila.avisos)

    def test_apellido_y_nombre_invierte(self):
        fila = leer(("Apellido y nombre",), ("Pérez Juan",)).filas[0]
        self.assertEqual((fila.datos["nombre"], fila.datos["apellido"]), ("Juan", "Pérez"))

    def test_apellido_y_nombre_con_coma(self):
        fila = leer(("Apellido y nombre",), ("Pérez, Juan",)).filas[0]
        self.assertEqual((fila.datos["nombre"], fila.datos["apellido"]), ("Juan", "Pérez"))

    def test_solo_nombre_se_parte_con_aviso(self):
        fila = leer(("Nombre",), ("Juan Pérez",)).filas[0]
        self.assertEqual((fila.datos["nombre"], fila.datos["apellido"]), ("Juan", "Pérez"))
        self.assertIn("Separé nombre y apellido: revisalo", fila.avisos)

    def test_solo_nombre_con_una_palabra_se_descarta(self):
        resultado = leer(("Nombre",), ("Juan",))
        self.assertEqual(resultado.filas, [])
        self.assertEqual(resultado.filas_invalidas[0].motivo, "Falta el apellido")

    def test_columnas_separadas_no_avisan(self):
        fila = leer(("Nombre", "Apellido"), ("Juan", "Pérez")).filas[0]
        self.assertEqual(fila.avisos, [])

    def test_apellido_vacio_con_nombre_de_varias_palabras_se_parte(self):
        fila = leer(("Nombre", "Apellido"), ("Juan Pérez", None)).filas[0]
        self.assertEqual((fila.datos["nombre"], fila.datos["apellido"]), ("Juan", "Pérez"))
        self.assertTrue(fila.avisos)

    def test_apellido_vacio_con_nombre_de_una_palabra_se_descarta(self):
        resultado = leer(("Nombre", "Apellido"), ("Juan", None))
        self.assertEqual(resultado.filas, [])
        self.assertEqual(resultado.filas_invalidas[0].motivo, "Falta el apellido")

    def test_nombre_vacio_se_descarta(self):
        resultado = leer(("Nombre", "Apellido", "Email"), (None, "Pérez", "a@b.com"))
        self.assertEqual(resultado.filas, [])
        self.assertEqual(resultado.filas_invalidas[0].motivo, "Falta el nombre")

    def test_fila_de_ejemplo_se_descarta(self):
        resultado = leer(
            ("Nombre", "Apellido", "Observaciones"),
            ("Ana", "Gómez", MARCA_FILA_EJEMPLO),
            ("Juan", "Pérez", "Una nota real"),
        )
        self.assertEqual([f.datos["nombre"] for f in resultado.filas], ["Juan"])
        self.assertEqual(resultado.filas_invalidas[0].fila_excel, 2)

    def test_las_observaciones_se_leen_tal_cual_sin_agregar_nada(self):
        fila = una_fila(
            ("Nombre", "Apellido", "Sexo", "Observaciones"),
            ("Ana", "Gómez", "X", "Lesión de rodilla\nrevisar"),
        )
        self.assertEqual(fila.datos["observaciones"], "Lesión de rodilla\nrevisar")


class EncabezadosTests(SimpleTestCase):
    def test_titulo_arriba_de_la_tabla(self):
        resultado = leer(
            ("Padrón 2026",),
            (None,),
            ("Nombre", "Apellido", "Celular"),
            ("Ana", "Gómez", 1122334455),
        )
        self.assertEqual(resultado.fila_encabezado, 3)
        self.assertEqual(resultado.filas[0].fila_excel, 4)

    def test_columnas_detectadas_e_ignoradas(self):
        resultado = leer(
            ("Nombre", "Apellido", "Celular", "Plan", "Monto", None),
            ("Ana", "Gómez", "1155", "Mensual", 15000, None),
        )
        self.assertEqual(
            resultado.columnas_detectadas,
            {"nombre": "Nombre", "apellido": "Apellido", "telefono": "Celular"},
        )
        self.assertEqual(resultado.columnas_ignoradas, ["Plan", "Monto"])

    def test_sin_columnas_de_nombre_levanta(self):
        with self.assertRaises(ColumnaRequeridaFaltante) as ctx:
            leer(("Plan", "Monto"), ("Mensual", 15000))
        self.assertEqual(ctx.exception.encabezados, ["Plan", "Monto"])

    def test_hoja_vacia_levanta(self):
        with self.assertRaises(ColumnaRequeridaFaltante):
            leer()

    def test_ningun_alias_es_de_una_letra(self):
        for campo, alias in ALIAS_ALUMNOS.items():
            for a in alias:
                with self.subTest(campo=campo, alias=a):
                    self.assertGreater(len(a), 1)

    def test_los_alias_son_normalizados(self):
        # `detectar_columnas` compara contra el encabezado ya normalizado:
        # un alias con mayúsculas o tildes jamás matchearía.
        for campo, alias in ALIAS_ALUMNOS.items():
            for a in alias:
                with self.subTest(campo=campo, alias=a):
                    self.assertEqual(a, parsing.normalizar_texto(a))

    def test_cada_alias_detecta_su_campo(self):
        # Fija también que no haya choques por prefijo entre alias de campos
        # distintos: cada alias, solo en una hoja con «Nombre» y «Apellido»,
        # tiene que caer en SU campo.
        for campo, alias in ALIAS_ALUMNOS.items():
            if campo in ("nombre", "apellido"):
                continue
            for a in alias:
                with self.subTest(campo=campo, alias=a):
                    if campo == "nombre_completo":
                        resultado = leer((a,), ("Juan Pérez",))
                    else:
                        resultado = leer(("Nombre", "Apellido", a), ("Juan", "Pérez", "x"))
                    self.assertEqual(resultado.columnas_detectadas.get(campo), a)

    def test_discapacidad_a_secas_es_el_tiene_y_el_detalle_gana_su_nombre(self):
        resultado = leer(
            ("Nombre", "Apellido", "Discapacidad", "Detalle discapacidad"),
            ("Juan", "Pérez", "Sí", "Sordera"),
        )
        self.assertEqual(resultado.columnas_detectadas["tiene_discapacidad"], "Discapacidad")
        self.assertEqual(
            resultado.columnas_detectadas["discapacidad_detalle"], "Detalle discapacidad"
        )

    def test_frecuencia_de_actividad_fisica_va_a_frecuencia_y_no_a_actividad(self):
        resultado = leer(
            ("Nombre", "Apellido", "Frecuencia de actividad física"), ("Juan", "Pérez", "diaria")
        )
        self.assertIn("frecuencia_actividad_previa", resultado.columnas_detectadas)
        self.assertNotIn("actividad_fisica_previa", resultado.columnas_detectadas)

    def test_celdas_combinadas(self):
        # El estado «Baja» está en una celda combinada de dos filas.
        resultado = leer(
            ("Nombre", "Apellido", "Estado"),
            ("Ana", "Gómez", "Baja"),
            ("Juan", "Pérez", None),
            merges=("C2:C3",),
        )
        self.assertEqual([f.datos["estado"] for f in resultado.filas], ["inactivo", "inactivo"])

    def test_telefono_numerico_pierde_el_punto_cero(self):
        fila = una_fila(("Nombre", "Apellido", "Celular"), ("Ana", "Gómez", 1122334455.0))
        self.assertEqual(fila.datos["telefono"], "1122334455")


class TopesTests(SimpleTestCase):
    def test_mas_del_tope_levanta_y_dice_mas_de_n(self):
        filas = [("Nombre", "Apellido")] + [
            (f"N{i}", f"A{i}") for i in range(MAX_FILAS_ALUMNOS + 1)
        ]
        with self.assertRaises(DemasiadasFilas) as ctx:
            leer(*filas)
        self.assertEqual(ctx.exception.maximo, MAX_FILAS_ALUMNOS)
        self.assertIn(f"más de {MAX_FILAS_ALUMNOS}", str(ctx.exception))

    def test_levanta_apenas_pasa_el_tope_sin_recorrer_el_resto(self):
        """Con 5.000 filas de datos, el lector no puede leer las 5.000 para
        rechazar después: corta en la fila `MAX_FILAS_ALUMNOS + 1`."""
        from unittest import mock

        filas = [("Nombre", "Apellido")] + [(f"N{i}", f"A{i}") for i in range(5000)]
        ws = hoja(*filas)
        leidas = set()
        real = lector._valor_celda

        def contar(ws_, fila, columna, merges):
            leidas.add(fila)
            return real(ws_, fila, columna, merges)

        with mock.patch.object(lector, "_valor_celda", side_effect=contar), \
                self.assertRaises(DemasiadasFilas):
            leer_hoja_alumnos(ws, hoy=HOY)
        self.assertLessEqual(max(leidas), MAX_FILAS_ALUMNOS + 2)

    def test_justo_el_tope_pasa(self):
        filas = [("Nombre", "Apellido")] + [
            (f"N{i}", f"A{i}") for i in range(MAX_FILAS_ALUMNOS)
        ]
        self.assertEqual(len(leer(*filas).filas), MAX_FILAS_ALUMNOS)

    def test_csv_con_demasiadas_filas_se_rechaza_sin_armar_el_libro_entero(self):
        texto = "Nombre,Apellido\n" + "".join(f"N{i},A{i}\n" for i in range(5000))
        archivo = io.BytesIO(texto.encode("utf-8"))
        archivo.name = "a.csv"
        with self.assertRaises(DemasiadasFilas):
            libro_desde_csv(archivo)

    def test_csv_justo_en_el_tope_se_arma(self):
        texto = "Nombre,Apellido\n" + "".join(
            f"N{i},A{i}\n" for i in range(MAX_FILAS_ALUMNOS)
        )
        archivo = io.BytesIO(texto.encode("utf-8"))
        archivo.name = "a.csv"
        self.assertEqual(
            len(parsear_archivo_alumnos(archivo, hoy=HOY).filas), MAX_FILAS_ALUMNOS
        )

    def test_doscientas_vacias_seguidas_cortan_con_aviso(self):
        filas = [("Nombre", "Apellido"), ("Ana", "Gómez")]
        filas += [(None, None)] * lector.FILAS_VACIAS_PARA_CORTAR
        filas += [("Juan", "Pérez")]
        resultado = leer(*filas)
        self.assertEqual([f.datos["nombre"] for f in resultado.filas], ["Ana"])
        self.assertTrue(resultado.avisos_lectura)
        self.assertIn("Dejé de leer en la fila", resultado.avisos_lectura[0])

    def test_vacias_hasta_el_final_no_avisan(self):
        filas = [("Nombre", "Apellido"), ("Ana", "Gómez")]
        filas += [(None, None)] * 10
        self.assertEqual(leer(*filas).avisos_lectura, [])

    def test_total_filas_leidas_cuenta_validas_e_invalidas(self):
        resultado = leer(
            ("Nombre", "Apellido"), ("Ana", "Gómez"), (None, "Solo"), (None, None), ("Juan", "Pérez")
        )
        self.assertEqual(resultado.total_filas_leidas, 3)
        self.assertEqual(len(resultado.filas), 2)
        self.assertEqual(len(resultado.filas_invalidas), 1)


class CsvTests(SimpleTestCase):
    FILAS = [
        ["Nombre", "Apellido", "Nacimiento", "Sexo", "Observaciones"],
        ["Ana", "Gómez", "03/05/1990", "F", "Vive lejos"],
        ["Juan", "Peña", "11/12/1985", "M", ""],
    ]

    def _csv(self, separador=";", codificacion="utf-8", bom=False):
        lineas = [separador.join(f) for f in self.FILAS]
        datos = "\r\n".join(lineas).encode(codificacion)
        return io.BytesIO((b"\xef\xbb\xbf" if bom else b"") + datos)

    def _leer(self, archivo):
        return leer_hoja_alumnos(libro_desde_csv(archivo).active, hoy=HOY)

    def _comparables(self, resultado):
        return [(f.fila_excel, f.datos, f.avisos, f.no_cargados) for f in resultado.filas]

    def test_separador_punto_y_coma_y_coma(self):
        for sep in (";", ","):
            with self.subTest(sep=repr(sep)):
                resultado = self._leer(self._csv(sep))
                self.assertEqual(len(resultado.filas), 2)
                self.assertEqual(resultado.filas[0].datos["fecha_nacimiento"], date(1990, 5, 3))

    def test_tabulador(self):
        self.assertEqual(len(self._leer(self._csv("\t")).filas), 2)

    def test_utf8_con_bom(self):
        resultado = self._leer(self._csv(bom=True))
        self.assertEqual(resultado.columnas_detectadas["nombre"], "Nombre")
        self.assertEqual(resultado.filas[1].datos["apellido"], "Peña")

    def test_cp1252_con_tildes(self):
        resultado = self._leer(self._csv(codificacion="cp1252"))
        self.assertEqual(resultado.filas[0].datos["apellido"], "Gómez")
        self.assertEqual(resultado.filas[1].datos["apellido"], "Peña")

    def test_mismo_resultado_que_el_xlsx_equivalente(self):
        desde_csv = self._leer(self._csv())
        desde_xlsx = leer(*[tuple(f) for f in self.FILAS])
        self.assertEqual(self._comparables(desde_csv), self._comparables(desde_xlsx))
        self.assertEqual(desde_csv.columnas_detectadas, desde_xlsx.columnas_detectadas)

    def test_utf16_levanta(self):
        crudo = "\r\n".join(";".join(f) for f in self.FILAS).encode("utf-16")
        with self.assertRaises(ArchivoCSVIlegible):
            libro_desde_csv(io.BytesIO(crudo))
        crudo = "\r\n".join(";".join(f) for f in self.FILAS).encode("utf-16-le")
        with self.assertRaises(ArchivoCSVIlegible):
            libro_desde_csv(io.BytesIO(crudo))

    def test_todo_llega_como_texto(self):
        csv_ = io.BytesIO("Nombre;Apellido;Celular\r\nAna;Gómez;1122334455".encode())
        fila = self._leer(csv_).filas[0]
        self.assertEqual(fila.datos["telefono"], "1122334455")

    def test_caracteres_de_control_no_rompen(self):
        csv_ = io.BytesIO("Nombre;Apellido\r\nAna\x07;Gómez".encode())
        self.assertEqual(self._leer(csv_).filas[0].datos["nombre"], "Ana")

    def test_parsear_archivo_alumnos_despacha_por_extension(self):
        archivo = self._csv()
        archivo.name = "padron.CSV"
        self.assertEqual(len(parsear_archivo_alumnos(archivo, hoy=HOY).filas), 2)

        wb = openpyxl.Workbook()
        for fila in self.FILAS:
            wb.active.append(fila)
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        buffer.name = "padron.xlsx"
        self.assertEqual(len(parsear_archivo_alumnos(buffer, hoy=HOY).filas), 2)

        # Sin `.name` se trata como xlsx.
        buffer.seek(0)
        del buffer.name
        self.assertEqual(len(parsear_archivo_alumnos(buffer, hoy=HOY).filas), 2)


class FachadaTests(SimpleTestCase):
    def test_exporta_lo_nuevo(self):
        for nombre in ("parsear_archivo_alumnos", "ALIAS_ALUMNOS", "leer_hoja_alumnos",
                       "libro_desde_csv", "MAX_FILAS_ALUMNOS", "MARCA_FILA_EJEMPLO",
                       "ArchivoCSVIlegible", "DemasiadasFilas"):
            with self.subTest(nombre=nombre):
                self.assertTrue(hasattr(parsing, nombre))

    def test_la_fachada_vieja_sigue_intacta(self):
        for nombre in ("parsear_archivo_plantillas", "parsear_archivo_biblioteca",
                       "leer_hoja_plantilla", "ColumnaRequeridaFaltante", "FilaInvalida",
                       "ALIAS_BIBLIOTECA", "normalizar_texto"):
            with self.subTest(nombre=nombre):
                self.assertTrue(hasattr(parsing, nombre))

    def test_el_lector_y_ciclo_son_django_free(self):
        import alumnos.ciclo as ciclo

        for modulo in (lector, ciclo):
            fuente = inspect.getsource(modulo)
            with self.subTest(modulo=modulo.__name__):
                self.assertNotIn("import django", fuente)
                self.assertNotIn("from django", fuente)
                self.assertNotIn("models", fuente.replace("models.py", ""))


class SincroniaConElModeloTests(SimpleTestCase):
    """Los valores que devuelve el lector tienen que ser choices válidas de
    `Alumno`: si alguien agrega o renombra una, esto rompe en vez de que el
    importador cree filas que el formulario después no puede editar."""

    def test_sexo_frecuencia_y_estado_son_choices_del_modelo(self):
        from alumnos.models import Alumno

        sexos = {lector.leer_sexo(v)[0] for v in (
            "masculino", "femenino", "prefiere no decir")}
        self.assertEqual(sexos, set(Alumno.Sexo.values))

        frecuencias = {lector.leer_frecuencia(v)[0] for v in (
            "diaria", "varias veces por semana", "una vez por semana", "ocasional")}
        self.assertEqual(frecuencias, set(Alumno.FrecuenciaActividad.values))

        estados = {lector.leer_estado(v)[0] for v in ("activo", "baja")}
        self.assertEqual(estados, set(Alumno.Estado.values))

    def test_los_datos_de_una_fila_son_campos_del_modelo(self):
        from alumnos.models import Alumno

        campos = {f.name for f in Alumno._meta.get_fields()}
        fila = una_fila(("Nombre", "Apellido"), ("Ana", "Gómez"))
        self.assertLessEqual(set(fila.datos), campos)


class RondaDeArreglos1Tests(SimpleTestCase):
    def test_negativos_en_tiene_no_inventan_datos_de_salud(self):
        for campo in ("Discapacidad", "Enfermedad crónica"):
            for valor in ("Ninguna", "ninguno", "No tiene", "No posee", "N/A", "—", "-",
                          "Ninguna conocida"):
                with self.subTest(campo=campo, valor=valor):
                    fila = una_fila(("Nombre", "Apellido", campo), ("Ana", "Gómez", valor))
                    d = fila.datos
                    self.assertFalse(d["tiene_discapacidad"] or d["tiene_enfermedad_cronica"])
                    self.assertEqual(d["discapacidad_detalle"], "")
                    self.assertEqual(d["enfermedad_cronica_detalle"], "")
                    self.assertEqual(fila.avisos, [])
                    self.assertEqual(fila.no_cargados, [])

    def test_otros_negativos_conocidos_tampoco_avisan(self):
        for valor in ("No presenta", "No refiere.", "Ningún problema", "no, ninguna", "Nada"):
            with self.subTest(valor=valor):
                fila = una_fila(("Nombre", "Apellido", "Discapacidad"), ("Ana", "Gómez", valor))
                self.assertFalse(fila.datos["tiene_discapacidad"])
                self.assertEqual(fila.datos["discapacidad_detalle"], "")
                self.assertEqual((fila.avisos, fila.no_cargados), ([], []))

    def test_un_texto_que_empieza_con_no_queda_en_no_con_aviso_y_no_se_pierde(self):
        """«No ve bien» ES el dato de salud: antes quedaba en falso sin aviso y
        el texto se perdía. Ni un sí inventado ni un no mudo: «no», aviso, y el
        texto original a observaciones."""
        for campo, clave, valor in (
            ("Discapacidad", "tiene_discapacidad", "No ve bien"),
            ("Discapacidad", "tiene_discapacidad", "No camina sin ayuda"),
            ("Enfermedad crónica", "tiene_enfermedad_cronica", "Ninguna salvo asma"),
        ):
            with self.subTest(valor=valor):
                fila = una_fila(("Nombre", "Apellido", campo), ("Ana", "Gómez", valor))
                self.assertFalse(fila.datos[clave])
                self.assertEqual(fila.datos[clave.replace("tiene_", "") + "_detalle"], "")
                self.assertEqual(len(fila.avisos), 1)
                self.assertIn(valor, fila.avisos[0])
                self.assertEqual(len(fila.no_cargados), 1)
                self.assertIn(f"«{valor}»", fila.no_cargados[0])

    def test_actividad_que_empieza_con_no_tambien_avisa(self):
        fila = una_fila(
            ("Nombre", "Apellido", "Actividad física previa"),
            ("Ana", "Gómez", "No desde la pandemia"),
        )
        self.assertFalse(fila.datos["actividad_fisica_previa"])
        self.assertTrue(fila.avisos)
        self.assertEqual(len(fila.no_cargados), 1)

    def test_una_palabra_que_empieza_con_no_no_es_un_no(self):
        # «Nota…» o «Nódulo…» no niegan nada: `\b` en el regex.
        self.assertFalse(lector.empieza_negando("Nódulo tiroideo"))
        fila = una_fila(("Nombre", "Apellido", "Enfermedad crónica"), ("Ana", "Gómez", "Nódulo tiroideo"))
        self.assertTrue(fila.datos["tiene_enfermedad_cronica"])

    def test_un_detalle_negativo_no_marca_tiene_como_si(self):
        fila = una_fila(
            ("Nombre", "Apellido", "Detalle discapacidad"), ("Ana", "Gómez", "No tiene")
        )
        self.assertFalse(fila.datos["tiene_discapacidad"])
        self.assertEqual(fila.avisos, [])

    def test_texto_libre_que_no_es_negativo_sigue_siendo_si(self):
        fila = una_fila(("Nombre", "Apellido", "Discapacidad"), ("Ana", "Gómez", "Motriz"))
        self.assertTrue(fila.datos["tiene_discapacidad"])
        self.assertEqual(fila.datos["discapacidad_detalle"], "Motriz")

    def test_lo_que_va_a_no_cargados_lleva_aviso_con_detalle_ya_cargado(self):
        for valor in ("Sí, visual", "Motriz"):
            with self.subTest(valor=valor):
                fila = una_fila(
                    ("Nombre", "Apellido", "Discapacidad", "Detalle discapacidad"),
                    ("Ana", "Gómez", valor, "Sordera"),
                )
                self.assertEqual(len(fila.no_cargados), 1)
                self.assertTrue(fila.avisos)

    def test_si_futbol_pasa_a_deportes_con_aviso(self):
        fila = una_fila(
            ("Nombre", "Apellido", "Actividad física previa"), ("Ana", "Gómez", "Sí, fútbol")
        )
        self.assertTrue(fila.datos["actividad_fisica_previa"])
        self.assertEqual(fila.datos["deportes_practica"], "fútbol")
        self.assertTrue(fila.avisos)
        self.assertEqual(fila.no_cargados, [])

    def test_si_futbol_con_deportes_ya_cargados_va_a_no_cargados_con_aviso(self):
        fila = una_fila(
            ("Nombre", "Apellido", "Actividad física previa", "Deportes"),
            ("Ana", "Gómez", "Sí, fútbol", "natación"),
        )
        self.assertEqual(fila.datos["deportes_practica"], "natación")
        self.assertEqual(len(fila.no_cargados), 1)
        self.assertTrue(fila.avisos)

    def test_nombre_completo_por_prefijo_no_se_pierde_en_silencio(self):
        resultado = leer(
            ("Nombre", "Apellido", "Alumno N°"), ("Ana", "Gómez", 17)
        )
        self.assertNotIn("nombre_completo", resultado.columnas_detectadas)
        self.assertEqual(resultado.columnas_ignoradas, ["Alumno N°"])

    def test_alumno_exacto_le_gana_a_nombre_y_apellido_del_tutor(self):
        resultado = leer(
            ("Alumno", "Nombre del tutor", "Apellido del tutor"),
            ("Juan Pérez", "Marta", "López"),
        )
        d = resultado.filas[0].datos
        self.assertEqual((d["nombre"], d["apellido"]), ("Juan", "Pérez"))
        self.assertEqual(
            resultado.columnas_ignoradas, ["Nombre del tutor", "Apellido del tutor"]
        )

    def test_alumno_y_apellido_importa_al_alumno_partido(self):
        resultado = leer(("Alumno", "Apellido"), ("Juan Pérez", "Pérez"))
        self.assertEqual(len(resultado.filas), 1)
        d = resultado.filas[0].datos
        self.assertEqual((d["nombre"], d["apellido"]), ("Juan", "Pérez"))

    def test_nombre_y_apellido_exactos_siguen_igual(self):
        resultado = leer(("Nombre", "Apellido", "Alumno"), ("Juan", "Pérez", "Otro Nombre"))
        d = resultado.filas[0].datos
        self.assertEqual((d["nombre"], d["apellido"]), ("Juan", "Pérez"))
        self.assertEqual(resultado.filas[0].avisos, [])

    def test_la_fila_del_ejemplo_no_decide_la_convencion_de_la_m(self):
        resultado = leer(
            ("Nombre", "Apellido", "Sexo", "Observaciones"),
            ("Ana", "Ejemplo", "F", MARCA_FILA_EJEMPLO),
            ("Juan", "Pérez", "M", None),
        )
        self.assertEqual(resultado.filas[0].datos["sexo"], "")
        self.assertEqual(resultado.filas[0].no_cargados, ["Sexo «M»"])


def _xlsx(filas, *, dimension="auto"):
    """Bytes de un xlsx con `filas`. `dimension`: "auto" (la que escribe
    openpyxl), None (sin <dimension>, como algunos generadores) o un rango
    («A1:B1048576», la dimensión inflada de una hoja con formato hasta el
    final)."""
    import re
    import zipfile

    wb = openpyxl.Workbook()
    ws = wb.active
    for fila in filas:
        ws.append(list(fila))
    buffer = io.BytesIO()
    wb.save(buffer)
    if dimension == "auto":
        return buffer.getvalue()
    reemplazo = b"" if dimension is None else f'<dimension ref="{dimension}"/>'.encode()
    salida = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as zin, \
            zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            datos = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                datos = re.sub(rb"<dimension[^>]*/>", reemplazo, datos)
            zout.writestr(item, datos)
    return salida.getvalue()


def _archivo_xlsx(filas, **kw):
    archivo = io.BytesIO(_xlsx(filas, **kw))
    archivo.name = "a.xlsx"
    return archivo


class VerificarFilasXlsxTests(SimpleTestCase):
    """La pasada `read_only` previa: un xlsx enorme se rechaza ANTES de la
    carga completa, que es lo que costaba ~50 s y ~800 MB con el worker
    único."""

    ENC = ("Nombre", "Apellido")

    def _filas(self, n):
        return [self.ENC] + [(f"N{i}", f"A{i}") for i in range(n)]

    def _parsear_sin_carga_completa(self, archivo):
        """Parsea fallando si se intenta la carga completa (no `read_only`)."""
        from unittest import mock

        real = openpyxl.load_workbook

        def solo_read_only(*a, **kw):
            if not kw.get("read_only"):
                raise AssertionError("cargó el libro entero")
            return real(*a, **kw)

        with mock.patch.object(parsing.openpyxl, "load_workbook", side_effect=solo_read_only):
            return parsear_archivo_alumnos(archivo, hoy=HOY)

    def test_muchas_filas_se_rechaza_sin_cargar_el_libro(self):
        filas = self._filas(lector.TOPE_MAX_ROW_SIN_VERIFICAR + 100)
        with self.assertRaises(DemasiadasFilas):
            self._parsear_sin_carga_completa(_archivo_xlsx(filas))

    def test_sin_dimension_tambien_se_rechaza_sin_cargar_el_libro(self):
        filas = self._filas(MAX_FILAS_ALUMNOS + 100)
        with self.assertRaises(DemasiadasFilas):
            self._parsear_sin_carga_completa(_archivo_xlsx(filas, dimension=None))

    def test_sin_dimension_y_pocas_filas_se_lee(self):
        resultado = parsear_archivo_alumnos(
            _archivo_xlsx(self._filas(10), dimension=None), hoy=HOY
        )
        self.assertEqual(len(resultado.filas), 10)

    def test_dimension_inflada_con_pocas_filas_se_lee(self):
        """`max_row` de un millón por formato aplicado hasta el final no son
        alumnos: no se rechaza por la dimensión."""
        archivo = _archivo_xlsx(self._filas(10), dimension="A1:B1048576")
        resultado = parsear_archivo_alumnos(archivo, hoy=HOY)
        self.assertEqual(len(resultado.filas), 10)

    def test_justo_el_tope_con_max_row_alto_se_lee(self):
        # Filas en blanco intercaladas inflan `max_row` sin sumar alumnos.
        filas = [self.ENC]
        for i in range(MAX_FILAS_ALUMNOS):
            filas += [(f"N{i}", f"A{i}"), (None, None), (None, None), (None, None)]
        resultado = parsear_archivo_alumnos(_archivo_xlsx(filas), hoy=HOY)
        self.assertEqual(len(resultado.filas), MAX_FILAS_ALUMNOS)

    def test_deja_el_archivo_rebobinado(self):
        archivo = _archivo_xlsx(self._filas(lector.TOPE_MAX_ROW_SIN_VERIFICAR + 100))
        with self.assertRaises(DemasiadasFilas):
            lector.verificar_filas_xlsx(archivo)
        self.assertEqual(archivo.tell(), 0)
        archivo = _archivo_xlsx(self._filas(5))
        lector.verificar_filas_xlsx(archivo)
        self.assertEqual(archivo.tell(), 0)
