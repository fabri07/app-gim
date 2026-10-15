"""Lector Django-free del archivo de ALUMNOS (`.xlsx` o `.csv`).

Mismo molde que `tabular.leer_hoja_biblioteca`: busca la fila de títulos,
detecta columnas por alias y recorre las filas. Sin ORM ni `timezone`: el
«hoy» lo pasa siempre quien llama.

**Regla de producto: no inventar ni perder en silencio.** Un dato opcional que
no se entiende queda vacío (o en su default) y la fila lleva un aviso; además
la fila lleva `no_cargados`, una lista de strings con el formato
`Campo «valor original»` (p. ej. `Sexo «X»`). **El lector NO escribe en
`observaciones`**: la Task 5 (services) concatena esos strings junto con lo que
ella misma detecta (email, teléfono, largos) y arma la línea
`Datos del archivo que no se pudieron cargar: …` al final de las observaciones.
Lo único que el lector hace con `observaciones` es leer la columna tal cual.

Se descarta una fila SOLO si falta nombre o apellido (y no se puede partir) o
si es una fila del archivo de ejemplo. Los largos de nombre/apellido, el email
y el teléfono se validan en services, no acá.
"""

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime

import openpyxl
from openpyxl.utils.datetime import from_excel

from alumnos.ciclo import rango_inicio_ciclo
from importaciones.parsing.comun import (
    ColumnaRequeridaFaltante,
    FilaInvalida,
    FILAS_BUSQUEDA_ENCABEZADO,
    _fila_vacia,
    _mapa_merges,
    _valor_celda,
    buscar_fila_encabezado,
    mejor_encabezado_parcial,
    normalizar_texto,
)

MAX_FILAS_ALUMNOS = 1000
FILAS_VACIAS_PARA_CORTAR = 200

# Más filas no vacías que esto en el archivo y no puede haber menos de
# `MAX_FILAS_ALUMNOS` alumnos: las únicas filas con datos que no son alumnos
# son las de arriba del encabezado (a lo sumo `FILAS_BUSQUEDA_ENCABEZADO`).
TOPE_FILAS_CON_DATOS = MAX_FILAS_ALUMNOS + FILAS_BUSQUEDA_ENCABEZADO

# Hasta este `max_row` la carga completa del xlsx es barata (medido el
# 2026-10-15: 3.000 filas de 16 columnas, 0,75 s y ~24 MB), así que no hace
# falta ninguna pasada previa: el lector rechaza solo. Por encima, se cuenta
# en modo `read_only` antes de cargar. El ×3 deja lugar a filas en
# blanco intercaladas y a filas con formato sin datos, que inflan `max_row`
# sin ser alumnos.
TOPE_MAX_ROW_SIN_VERIFICAR = MAX_FILAS_ALUMNOS * 3 + FILAS_BUSQUEDA_ENCABEZADO
MARCA_FILA_EJEMPLO = "Fila de ejemplo: borrala antes de subir"

# El orden importa: `detectar_columnas` resuelve el prefijo en este orden.
# `nombre_completo` antes que `nombre`; cada `*_detalle` antes que su `tiene_*`;
# `frecuencia` antes que `actividad`. Ningún alias es de una sola letra: una
# columna «M» o «F» sería un falso positivo en cualquier planilla.
ALIAS_ALUMNOS = {
    "nombre_completo": [
        "nombre y apellido", "nombre completo", "apellido y nombre",
        "apellidos y nombres", "nombres y apellidos", "alumno", "alumna",
    ],
    "nombre": ["nombre", "nombres", "first name"],
    "apellido": ["apellido", "apellidos", "last name", "surname"],
    "email": ["email", "e-mail", "e mail", "mail", "correo", "correo electronico"],
    "telefono": [
        "telefono", "tel", "celular", "cel", "movil", "whatsapp", "wsp",
        "numero de telefono",
    ],
    "fecha_nacimiento": [
        "fecha de nacimiento", "fecha nacimiento", "fecha_nacimiento",
        "nacimiento", "fecha de nac", "f nac", "fec nac", "cumpleanos",
    ],
    # «Fecha de alta» NO es alias: es el día que se anotó, quizá hace años.
    "inicio_ciclo": [
        "inicio del ciclo", "inicio ciclo", "fecha_inicio_ciclo",
        "inicio de ciclo", "inicio del plan", "inicio de plan",
        "fecha de inicio", "fecha inicio",
    ],
    "estado": ["estado", "situacion", "activo"],
    "sexo": ["sexo", "genero", "sex", "gender"],
    "frecuencia_actividad_previa": [
        "frecuencia_actividad_previa", "frecuencia actividad previa",
        "frecuencia de actividad", "frecuencia",
    ],
    "actividad_fisica_previa": [
        "actividad_fisica_previa", "actividad fisica previa",
        "hacia actividad fisica", "actividad fisica",
    ],
    "deportes_practica": [
        "deportes_practica", "deportes que practica", "deportes", "deporte",
    ],
    "discapacidad_detalle": [
        "discapacidad_detalle", "detalle discapacidad", "discapacidad detalle",
        "cual discapacidad", "tipo de discapacidad",
    ],
    "tiene_discapacidad": ["tiene_discapacidad", "tiene discapacidad", "discapacidad"],
    "enfermedad_cronica_detalle": [
        "enfermedad_cronica_detalle", "detalle enfermedad cronica",
        "enfermedad cronica detalle", "cual enfermedad",
    ],
    "tiene_enfermedad_cronica": [
        "tiene_enfermedad_cronica", "tiene enfermedad cronica",
        "enfermedad cronica", "enfermedades cronicas", "patologia", "patologias",
    ],
    "observaciones": [
        "observaciones", "observacion", "obs", "notas", "nota", "comentarios",
    ],
}

# Cómo se llama cada campo cuando hay que citarlo en un aviso o en
# `no_cargados`.
ETIQUETAS = {
    "fecha_nacimiento": "Fecha de nacimiento",
    "inicio_ciclo": "Inicio del ciclo",
    "estado": "Estado",
    "sexo": "Sexo",
    "frecuencia_actividad_previa": "Frecuencia de actividad previa",
    "actividad_fisica_previa": "Actividad física previa",
    "tiene_discapacidad": "Discapacidad",
    "tiene_enfermedad_cronica": "Enfermedad crónica",
}


class ArchivoCSVIlegible(Exception):
    """La codificación del CSV no es UTF-8 ni Windows-1252 (p. ej. UTF-16)."""


class DemasiadasFilas(Exception):
    """Más de `MAX_FILAS_ALUMNOS` filas con datos: se rechaza, no se trunca.

    No lleva el total: se levanta apenas se pasa el tope, sin recorrer el
    resto del archivo (un xlsx de 5 MB tardaba ~50 s y ~800 MB en llegar al
    final con el único worker de gunicorn). El mensaje dice «más de N»."""

    def __init__(self):
        self.maximo = MAX_FILAS_ALUMNOS
        super().__init__(
            f"El archivo tiene más de {MAX_FILAS_ALUMNOS} filas con datos"
        )


def verificar_filas_xlsx(archivo):
    """Rechaza con `DemasiadasFilas` un xlsx con demasiadas filas ANTES de la
    carga completa, que es lo que cuesta memoria y tiempo. Deja el archivo
    rebobinado.

    Primero mira `max_row` de la primera hoja en modo `read_only` (lee la
    dimensión declarada, casi gratis). Si no pasa de
    `TOPE_MAX_ROW_SIN_VERIFICAR`, no hay nada que verificar. Si la pasa, o si
    es `None` (el archivo no declara dimensión), NO se rechaza por eso: una
    dimensión inflada (formato aplicado hasta la fila un millón) no son
    alumnos, y sin dimensión no hay dato. En los dos casos se cuentan las
    filas con datos en streaming, con la misma regla de corte que el lector
    (200 vacías seguidas), y se rechaza recién al pasar
    `TOPE_FILAS_CON_DATOS`. El streaming de `read_only` no arma el libro en
    memoria y corta temprano. Medido el 2026-10-15 con un xlsx de 1,8 MB y
    26.000 filas: con dimensión, 0,25 s y ~1 MB; sin dimensión costaba 8,9 s
    y ~200 MB hasta que el lector cortaba, y con el conteo previo 2,3 s y
    ~3 MB (casi todo es openpyxl recorriendo la hoja al abrirla en
    `read_only` buscando la dimensión que no está).
    """
    try:
        wb = openpyxl.load_workbook(archivo, read_only=True, data_only=True)
        try:
            ws = wb[wb.sheetnames[0]]
            max_row = ws.max_row
            if max_row is not None and max_row <= TOPE_MAX_ROW_SIN_VERIFICAR:
                return
            con_datos = vacias_seguidas = 0
            for fila in ws.iter_rows(values_only=True):
                if _fila_vacia(list(fila)):
                    vacias_seguidas += 1
                    if con_datos and vacias_seguidas >= FILAS_VACIAS_PARA_CORTAR:
                        return
                    continue
                vacias_seguidas = 0
                con_datos += 1
                if con_datos > TOPE_FILAS_CON_DATOS:
                    raise DemasiadasFilas()
        finally:
            wb.close()
    finally:
        archivo.seek(0)


@dataclass(frozen=True)
class FilaAlumno:
    fila_excel: int
    # Claves = campos de `Alumno` + `fecha_inicio_ciclo`; fechas `date`.
    datos: dict
    avisos: list = field(default_factory=list)
    # `Campo «valor original»` de lo que no se pudo cargar (ver docstring).
    no_cargados: list = field(default_factory=list)


@dataclass(frozen=True)
class HojaAlumnos:
    fila_encabezado: int
    columnas_detectadas: dict   # campo -> texto del encabezado
    columnas_ignoradas: list    # encabezados sin campo
    advertencias_columnas: list
    filas: list = field(default_factory=list)            # [FilaAlumno]
    filas_invalidas: list = field(default_factory=list)  # [FilaInvalida]
    total_filas_leidas: int = 0
    # Avisos de la lectura entera (p. ej. «Dejé de leer en la fila X»).
    avisos_lectura: list = field(default_factory=list)


# --------------------------------------------------------------- textos ----

def _texto(valor):
    """Texto limpio de una celda: strip, espacios internos colapsados y un
    `float` sin decimales (`1122334455.0`) sin el `.0`."""
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return "Sí" if valor else "No"
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    if isinstance(valor, datetime):
        valor = valor.date()
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    return " ".join(str(valor).split())


# --------------------------------------------------------------- fechas ----

_RE_DMA = re.compile(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2}|\d{4})$")
_RE_ISO = re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})(?:[T ].*)?$")


def _como_numero(valor):
    """El número que representa una celda numérica o un texto solo de dígitos
    (en un CSV todo llega como texto), o `None`."""
    if isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return valor
    texto = _texto(valor)
    return int(texto) if texto.isdigit() else None


def leer_fecha(valor, *, hoy):
    """`(fecha | None, aviso | None)`. Vacío -> `(None, None)`; ilegible o
    inexistente -> `(None, aviso)`. Los textos son siempre día primero."""
    if valor is None or _texto(valor) == "":
        return None, None
    if isinstance(valor, datetime):
        return valor.date(), None
    if isinstance(valor, date):
        return valor, None

    numero = _como_numero(valor)
    if numero is not None:
        if 1900 <= numero <= 2100:
            return None, "Parece solo el año, no una fecha completa"
        if 1 <= numero <= 2958465:
            try:
                return from_excel(numero).date(), None
            except (ValueError, OverflowError, TypeError):
                pass
        return None, "No entendí la fecha"

    texto = _texto(valor)
    m = _RE_DMA.match(texto)
    if m:
        dia, mes, anio = int(m[1]), int(m[2]), int(m[3])
        if len(m[3]) == 2:
            anio += 1900 if anio > hoy.year % 100 else 2000
    else:
        m = _RE_ISO.match(texto)
        if not m:
            return None, "No entendí la fecha"
        anio, mes, dia = int(m[1]), int(m[2]), int(m[3])
    try:
        return date(anio, mes, dia), None
    except ValueError:
        return None, "La fecha no existe en el calendario"


def leer_fecha_nacimiento(valor, *, hoy):
    """`leer_fecha` + rechazo de lo anterior a 1900 o posterior a hoy."""
    fecha, aviso = leer_fecha(valor, hoy=hoy)
    if fecha is None:
        return None, aviso
    if fecha < date(1900, 1, 1) or fecha > hoy:
        return None, "La fecha de nacimiento está fuera de rango"
    return fecha, None


def leer_inicio_ciclo(valor, *, hoy):
    """`leer_fecha` + rango `[hoy - 27, hoy + 60]` (`alumnos.ciclo`)."""
    fecha, aviso = leer_fecha(valor, hoy=hoy)
    if fecha is None:
        return None, aviso
    desde, hasta = rango_inicio_ciclo(hoy)
    if not desde <= fecha <= hasta:
        return None, (
            f"El inicio del ciclo tiene que estar entre el "
            f"{desde:%d/%m/%Y} y el {hasta:%d/%m/%Y}"
        )
    return fecha, None


# --------------------------------------------------------------- estado ----

_ACTIVO = {"activo", "activa", "alta", "vigente", "si"}
_INACTIVO = {
    "inactivo", "inactiva", "baja", "de baja", "dado de baja",
    "dada de baja", "no",
}


def leer_estado(valor):
    """`("activo" | "inactivo", aviso | None)`. Vacío -> activo sin aviso;
    ilegible -> activo con aviso."""
    if isinstance(valor, bool):
        return ("activo" if valor else "inactivo"), None
    texto = normalizar_texto(_texto(valor))
    if texto == "":
        return "activo", None
    if texto in _ACTIVO:
        return "activo", None
    if texto in _INACTIVO:
        return "inactivo", None
    return "activo", "No entendí el estado: lo dejé como activo"


# ------------------------------------------------------------ booleanos ----

_VERDADERO = {"si", "s", "x", "1", "true", "verdadero", "yes", "y", "✓"}
_FALSO = {"no", "n", "0", "false", "falso", "", "-", "—", "–"}
_RE_SI_Y_RESTO = re.compile(r"^s[ií](?!\w)[\s,:;.\-–]*(.*)$", re.IGNORECASE)
# Detalles que dicen «no hay nada»: no infieren `tiene_* = True`.
_DETALLE_VACIO = {"", "no", "ninguna", "ninguno", "-", "n", "n/a", "na"}


def _es_negativo(norm):
    """«Ninguna», «N/A», «No tiene», «No posee», «Ninguna conocida»…: dicen que
    NO hay nada. Nunca pueden volverse un sí con el texto como detalle: eso
    sería inventar un dato de salud."""
    return (
        norm in _FALSO or norm in _DETALLE_VACIO
        or norm.startswith("no ") or norm.startswith("ningun")
    )


def leer_booleano(valor):
    """`(True | False | None, resto)`. `None` = texto que no es ni afirmativo ni
    negativo. «Sí, asma» -> `(True, "asma")`: el resto es lo que sobra después
    del sí. Los negativos («Ninguna», «No tiene», «—»…) dan `False`."""
    if isinstance(valor, bool):
        return valor, ""
    texto = _texto(valor)
    norm = normalizar_texto(texto)
    if norm in _VERDADERO:
        return True, ""
    if _es_negativo(norm):
        return False, ""
    m = _RE_SI_Y_RESTO.match(texto)
    if m:
        return True, m.group(1).strip()
    return None, ""


def detalle_informa(detalle):
    """¿El detalle dice algo (no es «ninguna», «-», vacío…)?"""
    return normalizar_texto(detalle) not in _DETALLE_VACIO


# ----------------------------------------------------------------- sexo ----

_SEXO_MASCULINO = {"masculino", "hombre", "varon", "masc"}
_SEXO_FEMENINO = {"femenino", "mujer", "fem"}
_SEXO_NO_DECIR = {"prefiere no decir", "no decir", "nd"}


def convencion_sexo(valores):
    """Qué significa la letra «M» en esta columna. Mirando la columna entera:
    con «F» y sin «H», M = masculino; con «H» y sin «F», M = femenino (mujer);
    con las dos o ninguna es ambigua (`None`)."""
    letras = {normalizar_texto(_texto(v)) for v in valores}
    hay_f, hay_h = "f" in letras, "h" in letras
    if hay_f and not hay_h:
        return "masculino"
    if hay_h and not hay_f:
        return "femenino"
    return None


def leer_sexo(valor, *, m_significa=None):
    """`(sexo | "", aviso | None)`. `m_significa` sale de `convencion_sexo`."""
    norm = normalizar_texto(_texto(valor))
    if norm == "":
        return "", None
    if norm in _SEXO_MASCULINO or norm == "h":
        return "masculino", None
    if norm in _SEXO_FEMENINO or norm == "f":
        return "femenino", None
    if norm in _SEXO_NO_DECIR:
        return "no_decir", None
    if norm == "m":
        if m_significa:
            return m_significa, None
        return "", "«M» es ambigua (¿masculino o mujer?): dejé el sexo vacío"
    return "", "No entendí el sexo"


# ----------------------------------------------------------- frecuencia ----

_FRECUENCIAS = {
    "diaria": "diaria", "diario": "diaria", "todos los dias": "diaria",
    "varias": "varias_por_semana", "varias por semana": "varias_por_semana",
    "varias veces por semana": "varias_por_semana",
    "una vez por semana": "una_por_semana",
    "una por semana": "una_por_semana",
    "semanal": "una_por_semana",
    "ocasional": "ocasional", "a veces": "ocasional",
    "esporadica": "ocasional", "esporadico": "ocasional",
    "ocasionalmente": "ocasional",
}
_RE_VECES = re.compile(
    r"^(\d+)(?:\s*(?:a|-|/|y)\s*(\d+))?\s*(?:veces|vez|x)?"
    r"(?:\s*(?:por|a la|x|/)\s*semana)?$"
)


def leer_frecuencia(valor):
    """`(frecuencia | "", aviso | None)`."""
    norm = normalizar_texto(_texto(valor)).replace("_", " ")
    if norm == "":
        return "", None
    if norm in _FRECUENCIAS:
        return _FRECUENCIAS[norm], None
    m = _RE_VECES.match(norm)
    if m:
        n = max(int(m[1]), int(m[2] or 0))
        if n == 1:
            return "una_por_semana", None
        if 2 <= n <= 6:
            return "varias_por_semana", None
        if n == 7:
            return "diaria", None
    return "", "No entendí la frecuencia de actividad"


# --------------------------------------------------------------- nombre ----

def partir_nombre_completo(texto, *, apellido_primero=False):
    """`(nombre, apellido)` de un nombre completo. Con coma: «Apellido,
    Nombre». Sin coma, la última palabra es el apellido y el resto el nombre
    (con `apellido_primero`, al revés: la primera palabra es el apellido). Una
    sola palabra queda como nombre y apellido vacío."""
    texto = " ".join(str(texto).split())
    if "," in texto:
        apellido, _, nombre = texto.partition(",")
        return nombre.strip(), apellido.strip()
    palabras = texto.split(" ")
    if len(palabras) < 2:
        return texto, ""
    if apellido_primero:
        return " ".join(palabras[1:]), palabras[0]
    return " ".join(palabras[:-1]), palabras[-1]


# ------------------------------------------------------------- una fila ----

def _par_tiene_detalle(v, campo_tiene, campo_detalle, avisos, no_cargados):
    """Resuelve `tiene_*` + `*_detalle` (spec «Booleanos» e «Implícitos»).
    Devuelve `(tiene, detalle)`."""
    etiqueta = ETIQUETAS[campo_tiene]
    crudo = v.get(campo_tiene)
    detalle = _texto(v.get(campo_detalle))
    tiene_vacio = _texto(crudo) == ""
    tiene = False

    if not tiene_vacio:
        valor, resto = leer_booleano(crudo)
        if valor is None:
            # Una sola columna «Discapacidad» con texto libre.
            tiene, resto = True, _texto(crudo)
        else:
            tiene = valor
        if resto:
            if detalle:
                avisos.append(
                    f"No pude pasar «{resto}» al detalle de {etiqueta.lower()}: "
                    "ya tenía uno cargado"
                )
                no_cargados.append(f"{etiqueta} «{_texto(crudo)}»")
            else:
                detalle = resto
                avisos.append(f"Pasé «{resto}» al detalle de {etiqueta.lower()}")

    if tiene_vacio and detalle_informa(detalle):
        tiene = True
        avisos.append(
            f"Marqué «{etiqueta.lower()}» como sí porque cargaste el detalle"
        )
    return tiene, detalle


def _construir_fila(fila_idx, v, *, modo_nombre, apellido_primero, m_significa, hoy):
    """`FilaAlumno` o `FilaInvalida` para una fila ya leída (`v`: campo ->
    valor crudo)."""
    observaciones = str(v.get("observaciones") or "").strip()
    if observaciones == MARCA_FILA_EJEMPLO:
        return FilaInvalida(fila_idx, "Es una fila de ejemplo")

    avisos, no_cargados = [], []

    if modo_nombre == "separado":
        nombre, apellido = _texto(v.get("nombre")), _texto(v.get("apellido"))
        partir = bool(nombre) and not apellido and len(nombre.split()) >= 2
        a_partir = nombre
    else:
        nombre = apellido = ""
        partir = True
        a_partir = _texto(v.get("nombre_completo"))
    if partir:
        nombre, apellido = partir_nombre_completo(
            a_partir, apellido_primero=apellido_primero
        )
        if nombre and apellido:
            avisos.append("Separé nombre y apellido: revisalo")
    if not nombre:
        return FilaInvalida(fila_idx, "Falta el nombre")
    if not apellido:
        return FilaInvalida(fila_idx, "Falta el apellido")

    def _no_leido(campo, aviso):
        avisos.append(aviso)
        no_cargados.append(f"{ETIQUETAS[campo]} «{_texto(v.get(campo))}»")

    nacimiento, aviso = leer_fecha_nacimiento(v.get("fecha_nacimiento"), hoy=hoy)
    if aviso:
        _no_leido("fecha_nacimiento", aviso)

    inicio, aviso = leer_inicio_ciclo(v.get("inicio_ciclo"), hoy=hoy)
    if aviso:
        _no_leido("inicio_ciclo", aviso)

    estado, aviso = leer_estado(v.get("estado"))
    if aviso:
        _no_leido("estado", aviso)

    sexo, aviso = leer_sexo(v.get("sexo"), m_significa=m_significa)
    if aviso:
        _no_leido("sexo", aviso)

    frecuencia, aviso = leer_frecuencia(v.get("frecuencia_actividad_previa"))
    if aviso:
        _no_leido("frecuencia_actividad_previa", aviso)

    deportes = _texto(v.get("deportes_practica"))
    crudo = v.get("actividad_fisica_previa")
    actividad_vacia = _texto(crudo) == ""
    actividad = False
    if not actividad_vacia:
        valor, resto = leer_booleano(crudo)
        if valor is None:
            _no_leido(
                "actividad_fisica_previa",
                "No entendí si hacía actividad física: lo dejé en «no»",
            )
        else:
            actividad = valor
            if resto and not deportes:
                deportes = resto
                avisos.append(f"Pasé «{resto}» a los deportes que practica")
            elif resto:
                avisos.append(
                    f"No pude pasar «{resto}» a los deportes: ya tenían algo cargado"
                )
                no_cargados.append(
                    f"{ETIQUETAS['actividad_fisica_previa']} «{_texto(crudo)}»"
                )
    if frecuencia and actividad_vacia:
        actividad = True
        avisos.append(
            "Marqué «actividad física previa» como sí porque cargaste la frecuencia"
        )

    tiene_disc, detalle_disc = _par_tiene_detalle(
        v, "tiene_discapacidad", "discapacidad_detalle", avisos, no_cargados
    )
    tiene_enf, detalle_enf = _par_tiene_detalle(
        v, "tiene_enfermedad_cronica", "enfermedad_cronica_detalle", avisos, no_cargados
    )

    datos = {
        "nombre": nombre,
        "apellido": apellido,
        "email": _texto(v.get("email")),
        "telefono": _texto(v.get("telefono")),
        "fecha_nacimiento": nacimiento,
        "sexo": sexo,
        "estado": estado,
        "fecha_inicio_ciclo": inicio,
        "actividad_fisica_previa": actividad,
        "frecuencia_actividad_previa": frecuencia,
        "deportes_practica": deportes,
        "tiene_discapacidad": tiene_disc,
        "discapacidad_detalle": detalle_disc,
        "tiene_enfermedad_cronica": tiene_enf,
        "enfermedad_cronica_detalle": detalle_enf,
        "observaciones": observaciones,
    }
    return FilaAlumno(fila_idx, datos, avisos, no_cargados)


# ------------------------------------------------------------- la hoja -----

def _encabezado_de_alumnos(ws):
    """Primero nombre + apellido; si no, una columna de nombre completo; si no,
    «Nombre» sola (se trata como nombre completo). `ColumnaRequeridaFaltante`
    si nada de eso."""
    for requeridos in (("nombre", "apellido"), ("nombre_completo",), ("nombre",)):
        encabezado = buscar_fila_encabezado(ws, ALIAS_ALUMNOS, requeridos)
        if encabezado is not None:
            return encabezado
    parcial = mejor_encabezado_parcial(ws, ALIAS_ALUMNOS)
    raise ColumnaRequeridaFaltante(
        "nombre y apellido", parcial.valores, fila=parcial.fila
    )


def leer_hoja_alumnos(ws, *, hoy):
    """Parsea una hoja de ALUMNOS. Devuelve un `HojaAlumnos`.

    Levanta `ColumnaRequeridaFaltante` sin columna de nombre y `DemasiadasFilas`
    con más de `MAX_FILAS_ALUMNOS` filas con datos. Fila inválida = se salta y
    se lista; nunca invalida la hoja."""
    encabezado = _encabezado_de_alumnos(ws)
    campos = dict(encabezado.campos)

    def exacto(campo):
        return normalizar_texto(encabezado.valores[campos[campo]]) in ALIAS_ALUMNOS[campo]

    separado = "nombre" in campos and "apellido" in campos
    # Una columna de nombre completo EXACTA le gana a un nombre/apellido
    # detectados solo por prefijo («Alumno» contra «Nombre del tutor»).
    if separado and "nombre_completo" in campos and exacto("nombre_completo") and not (
        exacto("nombre") and exacto("apellido")
    ):
        separado = False
    if separado:
        modo_nombre = "separado"
        # Un nombre completo detectado por prefijo («Alumno N°») no se usa:
        # que figure como columna ignorada y no como detectada.
        campos.pop("nombre_completo", None)
    else:
        modo_nombre = "completo"
        if "nombre_completo" not in campos:
            # «Nombre» sola: es el nombre completo.
            campos["nombre_completo"] = campos.pop("nombre")
        campos.pop("nombre", None)
        campos.pop("apellido", None)

    apellido_primero = modo_nombre == "completo" and normalizar_texto(
        encabezado.valores[campos["nombre_completo"]]
    ).startswith("apellido")

    mapa_merges = _mapa_merges(ws)
    ncols = len(encabezado.valores)
    ultima = ws.max_row or 0
    leidas = []        # (fila_idx, {campo: valor crudo})
    vacias_seguidas = 0
    avisos_lectura = []
    for fila_idx in range(encabezado.fila + 1, ultima + 1):
        valores = [_valor_celda(ws, fila_idx, c, mapa_merges) for c in range(1, ncols + 1)]
        if _fila_vacia(valores):
            vacias_seguidas += 1
            if vacias_seguidas >= FILAS_VACIAS_PARA_CORTAR:
                if fila_idx < ultima:
                    avisos_lectura.append(f"Dejé de leer en la fila {fila_idx}")
                break
            continue
        vacias_seguidas = 0
        leidas.append((fila_idx, {c: valores[i] for c, i in campos.items()}))
        # Adentro del loop: con un `max_row` de 150.000 no tiene sentido
        # recorrer el resto para después rechazar igual.
        if len(leidas) > MAX_FILAS_ALUMNOS:
            raise DemasiadasFilas()

    # La fila del ejemplo no cuenta: su «F» no puede decidir qué es la «M» de
    # los alumnos reales.
    m_significa = convencion_sexo(
        v.get("sexo") for _, v in leidas
        if str(v.get("observaciones") or "").strip() != MARCA_FILA_EJEMPLO
    )

    filas, invalidas = [], []
    for fila_idx, v in leidas:
        resultado = _construir_fila(
            fila_idx, v, modo_nombre=modo_nombre, apellido_primero=apellido_primero,
            m_significa=m_significa, hoy=hoy,
        )
        (invalidas if isinstance(resultado, FilaInvalida) else filas).append(resultado)

    usadas = set(campos.values())
    return HojaAlumnos(
        fila_encabezado=encabezado.fila,
        columnas_detectadas={c: str(encabezado.valores[i]) for c, i in campos.items()},
        columnas_ignoradas=[
            str(valor).strip() for i, valor in enumerate(encabezado.valores)
            if i not in usadas and valor is not None and str(valor).strip()
        ],
        advertencias_columnas=list(encabezado.advertencias),
        filas=filas,
        filas_invalidas=invalidas,
        total_filas_leidas=len(leidas),
        avisos_lectura=avisos_lectura,
    )


# ------------------------------------------------------------------ CSV ----

_ILEGALES = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def libro_desde_csv(archivo):
    """Un `openpyxl.Workbook` en memoria (una hoja) con el contenido del CSV,
    todo como texto, para pasarlo a `leer_hoja_alumnos`.

    Codificación `utf-8-sig` y, si falla, `cp1252`. Separador: `csv.Sniffer`
    entre `,`, `;` y tab; si falla (o elige uno que la primera línea no tiene),
    `;` si aparece en la primera línea y si no `,`. Otra codificación (UTF-16)
    levanta `ArchivoCSVIlegible`."""
    crudo = archivo.read()
    if hasattr(archivo, "seek"):
        archivo.seek(0)
    if isinstance(crudo, str):
        texto = crudo
    else:
        # cp1252 decodifica casi cualquier byte, incluido el UTF-16 (que trae
        # NUL): sin este corte «funcionaría» con basura.
        if crudo[:2] in (b"\xff\xfe", b"\xfe\xff") or b"\x00" in crudo:
            raise ArchivoCSVIlegible("El CSV no está en UTF-8 ni en Windows-1252")
        try:
            texto = crudo.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                texto = crudo.decode("cp1252")
            except UnicodeDecodeError:
                raise ArchivoCSVIlegible(
                    "El CSV no está en UTF-8 ni en Windows-1252"
                ) from None

    primera = next((linea for linea in texto.splitlines() if linea.strip()), "")
    try:
        separador = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t").delimiter
    except csv.Error:
        separador = None
    if separador is None or separador not in primera:
        separador = ";" if ";" in primera else ","

    wb = openpyxl.Workbook()
    ws = wb.active
    con_datos = 0
    for fila in csv.reader(io.StringIO(texto, newline=""), delimiter=separador):
        # Mismo tope que `verificar_filas_xlsx`: armar en memoria un libro de
        # decenas de miles de filas solo para rechazarlo después es justo el
        # costo que se quiere evitar con el worker único.
        if any(celda.strip() for celda in fila):
            con_datos += 1
            if con_datos > TOPE_FILAS_CON_DATOS:
                raise DemasiadasFilas()
        # openpyxl rechaza los caracteres de control al escribir la celda.
        ws.append([_ILEGALES.sub("", celda) for celda in fila])
    return wb
