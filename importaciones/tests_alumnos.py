"""Tests del servicio de previsualización de ALUMNOS (`importaciones/alumnos.py`).

El lector tiene sus propios tests; acá se prueba lo que agrega el servicio:
email/teléfono/largos, la marca de demo, los duplicados contra la base y dentro
del archivo, `nacen_vencidas`, el JSON y el costo en queries.
"""

import io
from datetime import timedelta
from unittest import mock

import openpyxl
from django.contrib.auth.models import User
from django.core.files.storage import default_storage
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.utils import timezone

from alumnos.models import Alumno
from importaciones.alumnos import (
    CAMPOS_IMPORTABLES,
    claves_de,
    confirmar_importacion_alumnos,
    previsualizar_importacion_alumnos,
)
from importaciones.models import Importacion
from importaciones.services import ImportacionInvalida
from tenants.models import Gimnasio

ENCABEZADOS = ("Nombre", "Apellido", "Email", "Celular")


def _archivo(filas, encabezados=ENCABEZADOS):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(list(encabezados))
    for fila in filas:
        ws.append(list(fila))
    buffer = io.BytesIO()
    wb.save(buffer)
    return SimpleUploadedFile(
        "alumnos.xlsx", buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


class BaseAlumnosTests(TestCase):
    def setUp(self):
        self.gimnasio = Gimnasio.objects.create(nombre="Gym", slug="gym")
        self.usuario = User.objects.create_user(username="staff", password="clave12345")

    def _previsualizar(self, filas, encabezados=ENCABEZADOS):
        importacion = previsualizar_importacion_alumnos(
            gimnasio=self.gimnasio,
            archivo=_archivo(filas, encabezados),
            usuario=self.usuario,
        )
        return importacion.resultado

    def _alumno(self, nombre, apellido, email="", telefono="", **extra):
        return Alumno.objects.create(
            gimnasio=self.gimnasio, nombre=nombre, apellido=apellido,
            email=email, telefono=telefono, **extra,
        )


class CreacionTests(BaseAlumnosTests):
    def test_crea_la_importacion_en_revision_con_el_archivo(self):
        importacion = previsualizar_importacion_alumnos(
            gimnasio=self.gimnasio,
            archivo=_archivo([("Juan", "Pérez", "", "")]),
            usuario=self.usuario,
        )
        self.assertEqual(importacion.tipo, Importacion.Tipo.ALUMNOS)
        self.assertEqual(importacion.estado, Importacion.Estado.EN_REVISION)
        self.assertEqual(importacion.creado_por, self.usuario)
        self.assertEqual(importacion.gimnasio, self.gimnasio)
        self.assertTrue(importacion.archivo.name)

    def test_forma_del_json(self):
        r = self._previsualizar([("Juan", "Pérez", "JUAN@x.com", "11 5555-1234")])
        self.assertEqual(r["version"], 1)
        for clave in (
            "fila_encabezado", "columnas_detectadas", "columnas_ignoradas",
            "advertencias_columnas", "avisos_lectura", "a_crear", "ya_existen",
            "filas_invalidas", "nacen_vencidas", "total_filas_leidas",
        ):
            self.assertIn(clave, r)
        fila = r["a_crear"][0]
        self.assertEqual(fila["fila_excel"], 2)
        self.assertEqual(fila["datos"]["email"], "juan@x.com")
        self.assertEqual(fila["claves"]["email"], "juan@x.com")
        self.assertEqual(fila["claves"]["nombre"], "juan perez")
        self.assertEqual(fila["claves"], claves_de(fila["datos"]))
        self.assertTrue(set(fila["datos"]) <= CAMPOS_IMPORTABLES)

    def test_fechas_van_como_iso_o_none(self):
        hoy = timezone.localdate()
        r = self._previsualizar(
            [("Juan", "Pérez", "", "", hoy - timedelta(days=3), None)],
            encabezados=ENCABEZADOS + ("Inicio del ciclo", "Fecha de nacimiento"),
        )
        datos = r["a_crear"][0]["datos"]
        self.assertEqual(datos["fecha_inicio_ciclo"], (hoy - timedelta(days=3)).isoformat())
        self.assertIsNone(datos["fecha_nacimiento"])

    def test_avisos_de_lectura_llegan_al_json(self):
        filas = [("Juan", "Pérez", "", "")] + [(None,) * 4] * 205
        filas.append(("Ana", "Gómez", "", ""))
        r = self._previsualizar(filas)
        self.assertTrue(r["avisos_lectura"])
        self.assertTrue(any("Dejé de leer" in a for a in r["avisos_lectura"]))

    def test_nombre_de_121_caracteres_se_descarta(self):
        r = self._previsualizar([("J" * 121, "Pérez", "", ""), ("Ana", "A" * 121, "", "")])
        self.assertEqual(r["a_crear"], [])
        self.assertEqual(len(r["filas_invalidas"]), 2)
        self.assertIn("120", r["filas_invalidas"][0]["motivo"])

    def test_nombre_de_120_caracteres_entra(self):
        r = self._previsualizar([("J" * 120, "Pérez", "", "")])
        self.assertEqual(len(r["a_crear"]), 1)


class ValoresTests(BaseAlumnosTests):
    def test_email_invalido_queda_vacio_con_aviso_y_observaciones(self):
        r = self._previsualizar([("Juan", "Pérez", "no-es-mail", "")])
        fila = r["a_crear"][0]
        self.assertEqual(fila["datos"]["email"], "")
        self.assertTrue(any("email" in a.lower() for a in fila["avisos"]))
        self.assertEqual(
            fila["datos"]["observaciones"],
            "Datos del archivo que no se pudieron cargar: Email «no-es-mail».",
        )

    def test_dos_emails_en_una_celda_es_invalido(self):
        r = self._previsualizar([("Juan", "Pérez", "a@x.com, b@x.com", "")])
        self.assertEqual(r["a_crear"][0]["datos"]["email"], "")

    def test_telefono_de_31_caracteres_queda_vacio(self):
        tel = "1" * 31
        r = self._previsualizar([("Juan", "Pérez", "", tel)])
        fila = r["a_crear"][0]
        self.assertEqual(fila["datos"]["telefono"], "")
        self.assertTrue(fila["avisos"])
        self.assertIn(f"Teléfono «{tel}»", fila["datos"]["observaciones"])
        self.assertIsNone(fila["claves"]["telefono"])

    def test_telefono_de_30_caracteres_entra(self):
        r = self._previsualizar([("Juan", "Pérez", "", "1" * 30)])
        self.assertEqual(r["a_crear"][0]["datos"]["telefono"], "1" * 30)

    def test_largo_de_opcional_desbordado_va_a_observaciones(self):
        largo = "x" * 201
        r = self._previsualizar(
            [("Juan", "Pérez", "", "", largo)],
            encabezados=ENCABEZADOS + ("Deportes",),
        )
        fila = r["a_crear"][0]
        self.assertEqual(fila["datos"]["deportes_practica"], "")
        self.assertIn(largo, fila["datos"]["observaciones"])

    def test_observaciones_existentes_y_linea_van_separadas_por_salto(self):
        r = self._previsualizar(
            [("Juan", "Pérez", "malo", "", "Lesión de rodilla")],
            encabezados=ENCABEZADOS + ("Observaciones",),
        )
        self.assertEqual(
            r["a_crear"][0]["datos"]["observaciones"],
            "Lesión de rodilla\nDatos del archivo que no se pudieron cargar: Email «malo».",
        )

    def test_junta_lo_del_lector_y_lo_propio_en_una_linea(self):
        r = self._previsualizar(
            [("Juan", "Pérez", "malo", "", "X")],
            encabezados=ENCABEZADOS + ("Sexo",),
        )
        obs = r["a_crear"][0]["datos"]["observaciones"]
        self.assertEqual(obs.count("Datos del archivo"), 1)
        self.assertIn("Sexo «X»", obs)
        self.assertIn("Email «malo»", obs)
        self.assertTrue(obs.endswith("."))

    def test_marca_demo_se_reescribe(self):
        r = self._previsualizar(
            [("Juan", "Pérez", "", "", "[demo]")],
            encabezados=ENCABEZADOS + ("Observaciones",),
        )
        fila = r["a_crear"][0]
        self.assertEqual(fila["datos"]["observaciones"], "[demo] (importado)")
        self.assertTrue(fila["avisos"])

    def test_fila_de_ejemplo_se_descarta(self):
        from importaciones.parsing import MARCA_FILA_EJEMPLO

        r = self._previsualizar(
            [("Juan", "Pérez", "", "", MARCA_FILA_EJEMPLO)],
            encabezados=ENCABEZADOS + ("Observaciones",),
        )
        self.assertEqual(r["a_crear"], [])
        self.assertEqual(len(r["filas_invalidas"]), 1)


class ErroresDeArchivoTests(BaseAlumnosTests):
    def test_cuenta_demo_se_rechaza_antes_de_leer(self):
        self.gimnasio.es_demo = True
        self.gimnasio.save(update_fields=["es_demo"])
        basura = SimpleUploadedFile("x.xlsx", b"no es un excel")
        with self.assertRaises(ImportacionInvalida) as ctx:
            previsualizar_importacion_alumnos(
                gimnasio=self.gimnasio, archivo=basura, usuario=self.usuario
            )
        self.assertIn("prueba", str(ctx.exception))
        self.assertEqual(Importacion.objects.count(), 0)

    def test_archivo_corrupto(self):
        with self.assertRaises(ImportacionInvalida):
            previsualizar_importacion_alumnos(
                gimnasio=self.gimnasio,
                archivo=SimpleUploadedFile("x.xlsx", b"no es un excel"),
                usuario=self.usuario,
            )

    def test_sin_columna_de_nombre(self):
        with self.assertRaises(ImportacionInvalida) as ctx:
            self._previsualizar([("a", "b")], encabezados=("Plan", "Monto"))
        self.assertIn("nombre", str(ctx.exception).lower())

    def test_csv_en_utf16_pide_guardarlo_de_otra_forma(self):
        archivo = SimpleUploadedFile("x.csv", "Nombre;Apellido\nJuan;Pérez".encode("utf-16"))
        with self.assertRaises(ImportacionInvalida) as ctx:
            previsualizar_importacion_alumnos(
                gimnasio=self.gimnasio, archivo=archivo, usuario=self.usuario
            )
        self.assertIn("UTF-8", str(ctx.exception))

    def test_demasiadas_filas(self):
        filas = [(f"N{i}", f"A{i}", "", "") for i in range(1001)]
        with self.assertRaises(ImportacionInvalida) as ctx:
            self._previsualizar(filas)
        self.assertIn("partes", str(ctx.exception))

    def test_si_no_se_puede_rebobinar_el_archivo_levanta(self):
        from unittest import mock

        hoja = previsualizar_importacion_alumnos(
            gimnasio=self.gimnasio, archivo=_archivo([("Juan", "Pérez", "", "")]),
            usuario=self.usuario,
        )
        self.assertTrue(hoja)

        class SinSeek(SimpleUploadedFile):
            def seek(self, *a, **k):
                raise OSError("no se puede")

        archivo = SinSeek("x.xlsx", _archivo([("Ana", "Gómez", "", "")]).read())
        from importaciones import alumnos as modulo

        real = modulo.parsear_archivo_alumnos
        with mock.patch.object(
            modulo, "parsear_archivo_alumnos",
            side_effect=lambda a, hoy: real(io.BytesIO(archivo.read()), hoy=hoy),
        ), self.assertLogs(
            "importaciones.alumnos", level="ERROR"
        ), self.assertRaises(ImportacionInvalida):
            previsualizar_importacion_alumnos(
                gimnasio=self.gimnasio, archivo=archivo, usuario=self.usuario
            )

    def test_csv_valido(self):
        archivo = SimpleUploadedFile(
            "x.csv", "Nombre;Apellido\nJuan;Pérez\n".encode("utf-8")
        )
        importacion = previsualizar_importacion_alumnos(
            gimnasio=self.gimnasio, archivo=archivo, usuario=self.usuario
        )
        self.assertEqual(len(importacion.resultado["a_crear"]), 1)


class DuplicadosTests(BaseAlumnosTests):
    def test_homonimo_con_email_invalido_se_crea_con_aviso(self):
        self._alumno("Juan", "Pérez", email="juan@x.com")
        r = self._previsualizar([("Juan", "Pérez", "no-es-mail", "")])
        self.assertEqual(r["ya_existen"], [])
        fila = r["a_crear"][0]
        self.assertTrue(any("mismo nombre" in a for a in fila["avisos"]))

    def test_mismo_email_en_mayusculas_y_mismo_nombre(self):
        existente = self._alumno("Juan", "Pérez", email="juan@x.com")
        r = self._previsualizar([("Juan", "Pérez", "JUAN@X.COM", "")])
        self.assertEqual(r["a_crear"], [])
        ya = r["ya_existen"][0]
        self.assertEqual(ya["origen"], "base")
        self.assertEqual(ya["alumno_id"], existente.pk)
        self.assertEqual((ya["nombre"], ya["apellido"]), ("Juan", "Pérez"))
        self.assertTrue(ya["motivo"])

    def test_telefono_con_y_sin_9(self):
        self._alumno("Juan", "Pérez", telefono="+54 9 11 5555-1234")
        r = self._previsualizar([("Juan", "Pérez", "", "011 5555-1234")])
        self.assertEqual(len(r["ya_existen"]), 1)

    def test_alumno_inactivo_tambien_cuenta(self):
        self._alumno("Juan", "Pérez", email="juan@x.com", estado=Alumno.Estado.INACTIVO)
        r = self._previsualizar([("Juan", "Pérez", "juan@x.com", "")])
        self.assertEqual(len(r["ya_existen"]), 1)

    def test_sin_contacto_el_nombre_basta(self):
        self._alumno("Juan", "Pérez", email="juan@x.com")
        r = self._previsualizar([("Pérez", "Juan", "", "")])  # nombre en otro orden
        self.assertEqual(len(r["ya_existen"]), 1)
        self.assertEqual(r["a_crear"], [])

    def test_hermano_con_mismo_telefono_y_otro_nombre_se_crea_con_aviso(self):
        self._alumno("Juan", "Pérez", telefono="1155551234")
        r = self._previsualizar([("Pedro", "Pérez", "", "1155551234")])
        self.assertEqual(r["ya_existen"], [])
        self.assertEqual(
            r["a_crear"][0]["avisos"], ["Comparte teléfono con Pérez, Juan"]
        )

    def test_mismo_email_otro_nombre_avisa_email(self):
        self._alumno("Juan", "Pérez", email="f@x.com")
        r = self._previsualizar([("Pedro", "Pérez", "f@x.com", "")])
        self.assertEqual(r["a_crear"][0]["avisos"], ["Comparte email con Pérez, Juan"])

    def test_mismo_nombre_otro_contacto_se_crea_con_aviso(self):
        self._alumno("Juan", "Pérez", email="juan@x.com")
        r = self._previsualizar([("Juan", "Pérez", "otro@x.com", "")])
        self.assertEqual(r["ya_existen"], [])
        self.assertEqual(
            r["a_crear"][0]["avisos"],
            ["Hay un alumno con el mismo nombre (Pérez, Juan): revisá que no sea el mismo"],
        )

    def test_aislamiento_otro_gimnasio_no_marca_duplicado(self):
        otro = Gimnasio.objects.create(nombre="Otro", slug="otro")
        Alumno.objects.create(
            gimnasio=otro, nombre="Juan", apellido="Pérez", email="juan@x.com"
        )
        r = self._previsualizar([("Juan", "Pérez", "juan@x.com", "")])
        self.assertEqual(r["ya_existen"], [])
        self.assertEqual(len(r["a_crear"]), 1)
        self.assertEqual(r["a_crear"][0]["avisos"], [])

    def test_duplicado_dentro_del_archivo_gana_la_primera(self):
        r = self._previsualizar([
            ("Juan", "Pérez", "juan@x.com", ""),
            ("Juan", "Pérez", "JUAN@x.com", ""),
        ])
        self.assertEqual(len(r["a_crear"]), 1)
        ya = r["ya_existen"][0]
        self.assertEqual(ya["origen"], "archivo")
        self.assertEqual(ya["fila_excel"], 3)
        self.assertIsNone(ya["alumno_id"])
        self.assertIn("fila 2", ya["motivo"])

    def test_dentro_del_archivo_nombre_sin_contacto(self):
        r = self._previsualizar([("Juan", "Pérez", "", ""), ("Juan", "Pérez", "", "")])
        self.assertEqual(len(r["a_crear"]), 1)
        self.assertEqual(len(r["ya_existen"]), 1)

    def test_dentro_del_archivo_hermano_se_crea_con_aviso(self):
        r = self._previsualizar([
            ("Juan", "Pérez", "", "1155551234"),
            ("Pedro", "Pérez", "", "1155551234"),
        ])
        self.assertEqual(len(r["a_crear"]), 2)
        self.assertEqual(r["a_crear"][1]["avisos"], ["Comparte teléfono con Pérez, Juan"])

    # «Sin contacto ⇒ comparar por nombre» mira LOS DOS lados. El caso real:
    # el gimnasio cargó a mano solo el nombre y después importa el padrón con
    # celulares; crearlo de nuevo duplica a la persona y le cobra dos veces.
    def test_existente_sin_contacto_y_fila_con_contacto_es_duplicado(self):
        existente = self._alumno("Juan", "Pérez")
        r = self._previsualizar([("Juan", "Pérez", "juan@x.com", "1155551234")])
        self.assertEqual(r["a_crear"], [])
        self.assertEqual(r["ya_existen"][0]["alumno_id"], existente.pk)

    def test_existente_sin_contacto_y_otro_nombre_se_crea(self):
        self._alumno("Juan", "Pérez")
        r = self._previsualizar([("Pedro", "Pérez", "pedro@x.com", "")])
        self.assertEqual(r["ya_existen"], [])
        self.assertEqual(len(r["a_crear"]), 1)

    def test_dentro_del_archivo_primera_sin_contacto_segunda_con_contacto(self):
        r = self._previsualizar([
            ("Juan", "Pérez", "", ""),
            ("Juan", "Pérez", "juan@x.com", ""),
        ])
        self.assertEqual(len(r["a_crear"]), 1)
        self.assertEqual(r["ya_existen"][0]["origen"], "archivo")
        self.assertIn("fila 2", r["ya_existen"][0]["motivo"])

    def test_con_contacto_en_los_dos_lados_sigue_haciendo_falta_el_contacto(self):
        # Que un lado sin contacto cuente por nombre no tiene que tragarse al
        # homónimo con OTRO contacto, aunque además exista uno sin contacto
        # con otro nombre: la regla de los dos lados con contacto no cambia.
        self._alumno("Juan", "Pérez", email="juan@x.com")
        self._alumno("Pedro", "Pérez")
        r = self._previsualizar([("Juan", "Pérez", "otro@x.com", "")])
        self.assertEqual(r["ya_existen"], [])
        self.assertEqual(len(r["a_crear"]), 1)


class NacenVencidasTests(BaseAlumnosTests):
    def _con_inicio(self, dias_atras):
        hoy = timezone.localdate()
        return self._previsualizar(
            [("Juan", "Pérez", "", "", hoy - timedelta(days=dias_atras)),
             ("Ana", "Gómez", "", "", None)],
            encabezados=ENCABEZADOS + ("Inicio del ciclo",),
        )

    def _tolerancia(self, dias, activacion="ayer"):
        # `update()` y no `save()`: la señal estampa la activación en hoy al
        # prender la tolerancia, y acá hace falta fijarla a mano.
        Gimnasio.objects.filter(pk=self.gimnasio.pk).update(
            dias_tolerancia_pago=dias,
            fecha_activacion_bloqueo=(
                timezone.localdate() - timedelta(days=400)
                if activacion == "ayer" else activacion
            ),
        )
        self.gimnasio.refresh_from_db()

    def test_sin_tolerancia_es_cero(self):
        self.assertEqual(self._con_inicio(20)["nacen_vencidas"], 0)

    def test_borde_exacto_cuenta(self):
        # `pagos.acceso` bloquea con periodo_inicio <= hoy - tolerancia.
        self._tolerancia(10)
        self.assertEqual(self._con_inicio(20)["nacen_vencidas"], 1)
        self.assertEqual(self._con_inicio(10)["nacen_vencidas"], 1)
        self.assertEqual(self._con_inicio(9)["nacen_vencidas"], 0)

    def test_tolerancia_cero_con_inicio_hoy_cuenta(self):
        self._tolerancia(0)
        self.assertEqual(self._con_inicio(0)["nacen_vencidas"], 1)

    def test_sin_fecha_de_activacion_no_cuenta(self):
        self._tolerancia(10, activacion=None)
        self.assertEqual(self._con_inicio(20)["nacen_vencidas"], 0)

    def test_activacion_posterior_al_periodo_no_cuenta(self):
        self._tolerancia(10, activacion=timezone.localdate() - timedelta(days=5))
        self.assertEqual(self._con_inicio(20)["nacen_vencidas"], 0)

    def test_inicio_futuro_no_cuenta(self):
        self._tolerancia(0)
        self.assertEqual(self._con_inicio(-5)["nacen_vencidas"], 0)

    def test_un_inactivo_no_cuenta(self):
        # El cron no le emite cuotas a un inactivo: nunca bloquea.
        self._tolerancia(10)
        hoy = timezone.localdate()
        r = self._previsualizar(
            [("Juan", "Pérez", "", "", hoy - timedelta(days=20), "Inactivo"),
             ("Ana", "Gómez", "", "", hoy - timedelta(days=20), "Activo")],
            encabezados=ENCABEZADOS + ("Inicio del ciclo", "Estado"),
        )
        self.assertEqual(r["nacen_vencidas"], 1)


class NaceBloqueadaTests(TestCase):
    def test_usa_el_periodo_real_del_cron_no_el_ancla(self):
        from pagos.acceso import nace_bloqueada

        hoy = timezone.localdate()
        gimnasio = Gimnasio.objects.create(nombre="G", slug="g")
        Gimnasio.objects.filter(pk=gimnasio.pk).update(
            dias_tolerancia_pago=10,
            fecha_activacion_bloqueo=hoy - timedelta(days=400),
        )
        gimnasio.refresh_from_db()
        # Ancla de hace 30 días: el ciclo vigente arrancó hace 2, no bloquea.
        self.assertFalse(nace_bloqueada(gimnasio, hoy - timedelta(days=30), hoy))
        self.assertTrue(nace_bloqueada(gimnasio, hoy - timedelta(days=10), hoy))
        self.assertFalse(nace_bloqueada(gimnasio, None, hoy))


class EscalaTests(BaseAlumnosTests):
    def _queries(self, filas):
        with CaptureQueriesContext(connection) as consultas:
            previsualizar_importacion_alumnos(
                gimnasio=self.gimnasio, archivo=_archivo(filas), usuario=self.usuario
            )
        return len(consultas)

    def test_costo_no_crece_con_las_filas(self):
        def filas(n, desde=0):
            return [
                (f"Nom{i}", f"Ape{i}", f"a{i}@x.com", f"11555{i:05d}")
                for i in range(desde, desde + n)
            ]

        chico = self._queries(filas(20))
        grande = self._queries(filas(200, 1000))
        self.assertLessEqual(abs(grande - chico), 2)

    def test_costo_no_crece_con_los_alumnos_existentes(self):
        filas = [(f"Nom{i}", f"Ape{i}", f"a{i}@x.com", "") for i in range(20)]
        for i in range(5):
            self._alumno(f"Ex{i}", f"Ist{i}", email=f"e{i}@x.com")
        pocos = self._queries(filas)
        for i in range(5, 50):
            self._alumno(f"Ex{i}", f"Ist{i}", email=f"e{i}@x.com")
        muchos = self._queries(filas)
        self.assertLessEqual(abs(muchos - pocos), 2)


class ConfirmarTests(BaseAlumnosTests):
    def _importacion(self, filas, encabezados=ENCABEZADOS, gimnasio=None):
        return previsualizar_importacion_alumnos(
            gimnasio=gimnasio or self.gimnasio,
            archivo=_archivo(filas, encabezados),
            usuario=self.usuario,
        )

    def _confirmar(self, importacion, gimnasio=None):
        with self.captureOnCommitCallbacks(execute=True):
            return confirmar_importacion_alumnos(
                importacion=importacion, gimnasio=gimnasio or self.gimnasio
            )

    def test_crea_activos_sin_ancla_ni_baja_y_marcados_importados(self):
        imp = self._importacion([("Juan", "Pérez", "J@x.com", "11 5555-1234")])
        resumen = self._confirmar(imp)
        a = Alumno.objects.get(gimnasio=self.gimnasio)
        self.assertEqual((a.nombre, a.apellido, a.email), ("Juan", "Pérez", "j@x.com"))
        self.assertEqual(a.estado, Alumno.Estado.ACTIVO)
        self.assertIsNone(a.fecha_inicio_ciclo)
        self.assertIsNone(a.fecha_baja)
        self.assertTrue(a.importado)
        self.assertEqual(resumen["creados"], 1)
        self.assertEqual(resumen["sin_fecha_inicio"], 1)

    def test_existente_cargado_despues_del_preview_sin_contacto_no_se_duplica(self):
        """La confirmación re-evalúa contra la base con la MISMA regla: si
        entre el preview y el confirmar alguien cargó a mano a «Juan Pérez»
        solo con el nombre, la fila con contacto no lo duplica."""
        imp = self._importacion([("Juan", "Pérez", "juan@x.com", "1155551234")])
        self.assertEqual(len(imp.resultado["a_crear"]), 1)
        self._alumno("Juan", "Pérez")
        resumen = self._confirmar(imp)
        self.assertEqual(resumen["creados"], 0)
        self.assertEqual(Alumno.objects.filter(gimnasio=self.gimnasio).count(), 1)

    def test_respeta_columnas_ancla_y_estado_del_archivo(self):
        hoy = timezone.localdate()
        ancla = hoy - timedelta(days=3)
        imp = self._importacion(
            [("Juan", "Pérez", "", "", ancla, "Inactivo", "1990-05-03")],
            encabezados=ENCABEZADOS + ("Inicio del ciclo", "Estado", "Fecha de nacimiento"),
        )
        resumen = self._confirmar(imp)
        a = Alumno.objects.get(gimnasio=self.gimnasio)
        self.assertEqual(a.fecha_inicio_ciclo, ancla)
        self.assertEqual(a.estado, Alumno.Estado.INACTIVO)
        self.assertEqual(a.fecha_nacimiento.isoformat(), "1990-05-03")
        self.assertIsNone(a.fecha_baja)
        self.assertEqual(resumen["sin_fecha_inicio"], 0)

    def test_un_inactivo_importado_no_cuenta_en_bajas_del_mes(self):
        from tenants.analitica import altas_y_bajas_por_mes

        imp = self._importacion(
            [("Juan", "Pérez", "", "", "Inactivo")],
            encabezados=ENCABEZADOS + ("Estado",),
        )
        self._confirmar(imp)
        filas = altas_y_bajas_por_mes(self.gimnasio, hoy=timezone.localdate())
        self.assertEqual(sum(f["bajas"] for f in filas), 0)
        self.assertEqual(sum(f["altas"] for f in filas), 0)

    def test_sin_ancla_el_cron_no_emite_cuotas_ni_bloquea(self):
        from pagos.acceso import bloqueo_de
        from pagos.models import Cuota, generar_pagos_pendientes

        Gimnasio.objects.filter(pk=self.gimnasio.pk).update(
            dias_tolerancia_pago=0,
            fecha_activacion_bloqueo=timezone.localdate() - timedelta(days=400),
        )
        self.gimnasio.refresh_from_db()
        self._confirmar(self._importacion([("Juan", "Pérez", "", "")]))
        hoy = timezone.localdate()
        generar_pagos_pendientes(hoy)
        self.assertEqual(Cuota.objects.filter(gimnasio=self.gimnasio).count(), 0)
        a = Alumno.objects.get(gimnasio=self.gimnasio)
        self.assertIsNone(bloqueo_de(a, hoy=hoy))

    def test_gimnasio_no_mixto_completa_el_sexo(self):
        Gimnasio.objects.filter(pk=self.gimnasio.pk).update(tipo_publico="mujeres")
        self.gimnasio.refresh_from_db()
        self._confirmar(self._importacion([("Ana", "Gómez", "", "")]))
        a = Alumno.objects.get(gimnasio=self.gimnasio)
        self.assertEqual(a.sexo, self.gimnasio.sexo_implicito())
        self.assertTrue(a.sexo)

    def test_gimnasio_mixto_deja_el_sexo_vacio(self):
        self._confirmar(self._importacion([("Ana", "Gómez", "", "")]))
        self.assertEqual(Alumno.objects.get(gimnasio=self.gimnasio).sexo, "")

    def test_segunda_confirmacion_levanta(self):
        imp = self._importacion([("Juan", "Pérez", "", "")])
        self._confirmar(imp)
        with self.assertRaisesMessage(ImportacionInvalida, "Esta importación ya fue procesada."):
            confirmar_importacion_alumnos(importacion=imp, gimnasio=self.gimnasio)
        self.assertEqual(Alumno.objects.filter(gimnasio=self.gimnasio).count(), 1)

    def test_importacion_de_otro_gimnasio_se_rechaza(self):
        otro = Gimnasio.objects.create(nombre="Otro", slug="otro")
        imp = self._importacion([("Juan", "Pérez", "", "")], gimnasio=otro)
        with self.assertRaises(ImportacionInvalida):
            confirmar_importacion_alumnos(importacion=imp, gimnasio=self.gimnasio)
        self.assertEqual(Alumno.objects.count(), 0)

    def test_importacion_de_otro_tipo_se_rechaza(self):
        imp = self._importacion([("Juan", "Pérez", "", "")])
        Importacion.objects.filter(pk=imp.pk).update(tipo=Importacion.Tipo.BIBLIOTECA)
        imp.refresh_from_db()
        with self.assertRaises(ImportacionInvalida):
            confirmar_importacion_alumnos(importacion=imp, gimnasio=self.gimnasio)

    def test_alumno_creado_entre_preview_y_confirmar_se_saltea(self):
        imp = self._importacion([
            ("Juan", "Pérez", "j@x.com", ""), ("Ana", "Gómez", "a@x.com", ""),
        ])
        self._alumno("Juan", "Pérez", email="j@x.com")
        resumen = self._confirmar(imp)
        self.assertEqual(resumen["creados"], 1)
        self.assertEqual(resumen["salteados_al_confirmar"], 1)
        self.assertEqual(Alumno.objects.filter(gimnasio=self.gimnasio).count(), 2)

    def test_resumen_cuenta_ya_existian_y_descartadas(self):
        self._alumno("Juan", "Pérez", email="j@x.com")
        imp = self._importacion([
            ("Juan", "Pérez", "j@x.com", ""), ("Ana", "Gómez", "", ""),
            ("", "SinNombre", "", ""),
        ])
        r = imp.resultado
        resumen = self._confirmar(imp)
        self.assertEqual(resumen["ya_existian"], len(r["ya_existen"]))
        self.assertEqual(resumen["descartadas"], len(r["filas_invalidas"]))
        self.assertEqual(resumen["creados"], 1)
        self.assertEqual(resumen["salteados_al_confirmar"], 0)

    def test_resultado_achicado_estado_y_archivo_borrado(self):
        imp = self._importacion([("Juan", "Pérez", "", "")])
        nombre = imp.archivo.name
        self.assertTrue(nombre)
        self.assertTrue(default_storage.exists(nombre))
        self._confirmar(imp)
        imp.refresh_from_db()
        self.assertEqual(imp.estado, Importacion.Estado.CONFIRMADA)
        self.assertIsNotNone(imp.confirmado_en)
        self.assertEqual(set(imp.resultado), {"version", "resumen"})
        self.assertEqual(
            set(imp.resultado["resumen"]),
            {"creados", "ya_existian", "descartadas", "salteados_al_confirmar"},
        )
        self.assertFalse(imp.archivo)
        self.assertFalse(default_storage.exists(nombre))

    def test_falla_del_storage_no_propaga(self):
        imp = self._importacion([("Juan", "Pérez", "", "")])
        with mock.patch(
            "django.db.models.fields.files.FieldFile.delete",
            side_effect=OSError("storage caído"),
        ), self.assertLogs("importaciones.alumnos", level="WARNING"):
            self._confirmar(imp)
        imp.refresh_from_db()
        self.assertEqual(imp.estado, Importacion.Estado.CONFIRMADA)
        self.assertEqual(Alumno.objects.filter(gimnasio=self.gimnasio).count(), 1)

    def test_no_se_duplican_entre_si_al_confirmar(self):
        imp = self._importacion([("Juan", "Pérez", "", ""), ("Ana", "Gómez", "", "")])
        # Se fuerza una fila repetida en el JSON (como si el preview fuera viejo).
        r = imp.resultado
        r["a_crear"].append(dict(r["a_crear"][0], fila_excel=99))
        Importacion.objects.filter(pk=imp.pk).update(resultado=r)
        imp.refresh_from_db()
        resumen = self._confirmar(imp)
        self.assertEqual(resumen["creados"], 2)
        self.assertEqual(resumen["salteados_al_confirmar"], 1)

    def test_costo_no_crece_con_las_filas(self):
        def confirmar(filas):
            imp = self._importacion(filas)
            with CaptureQueriesContext(connection) as consultas:
                with self.captureOnCommitCallbacks(execute=True):
                    confirmar_importacion_alumnos(importacion=imp, gimnasio=self.gimnasio)
            return len(consultas)

        def filas(n, desde=0):
            return [
                (f"Nom{i}", f"Ape{i}", f"a{i}@x.com", f"11555{i:05d}")
                for i in range(desde, desde + n)
            ]

        chico = confirmar(filas(20))
        grande = confirmar(filas(200, 1000))
        self.assertLessEqual(grande, chico + 5)


class ConfirmarContactoInvalidoTests(BaseAlumnosTests):
    """Un homónimo con contacto inválido se crea en el preview; al confirmar
    tiene que seguir creándose (el contacto se mide ANTES de validar)."""

    def _importar(self, filas):
        imp = previsualizar_importacion_alumnos(
            gimnasio=self.gimnasio, archivo=_archivo(filas), usuario=self.usuario
        )
        prometidos = len(imp.resultado["a_crear"])
        with self.captureOnCommitCallbacks(execute=True):
            resumen = confirmar_importacion_alumnos(
                importacion=imp, gimnasio=self.gimnasio
            )
        return prometidos, resumen

    def test_email_invalido_con_homonimo_en_la_base(self):
        self._alumno("Juan", "Pérez", email="j@x.com")
        prometidos, resumen = self._importar([("Juan", "Pérez", "no-es-email", "")])
        self.assertEqual((prometidos, resumen["creados"]), (1, 1))
        self.assertEqual(resumen["salteados_al_confirmar"], 0)

    def test_telefono_ilegible_con_homonimo_en_la_base(self):
        self._alumno("Juan", "Pérez", email="j@x.com")
        prometidos, resumen = self._importar([("Juan", "Pérez", "", "abc")])
        self.assertEqual((prometidos, resumen["creados"]), (1, 1))

    def test_dos_homonimos_con_contacto_invalido_en_el_archivo(self):
        prometidos, resumen = self._importar([
            ("Juan", "Pérez", "no-es-email", ""), ("Juan", "Pérez", "", "abc"),
        ])
        self.assertEqual((prometidos, resumen["creados"]), (2, 2))

    def test_importacion_vieja_sin_trae_contacto_sigue_confirmando(self):
        imp = previsualizar_importacion_alumnos(
            gimnasio=self.gimnasio,
            archivo=_archivo([("Juan", "Pérez", "j@x.com", "")]),
            usuario=self.usuario,
        )
        r = imp.resultado
        for fila in r["a_crear"]:
            del fila["trae_contacto"]
        Importacion.objects.filter(pk=imp.pk).update(resultado=r)
        imp.refresh_from_db()
        with self.captureOnCommitCallbacks(execute=True):
            resumen = confirmar_importacion_alumnos(
                importacion=imp, gimnasio=self.gimnasio
            )
        self.assertEqual(resumen["creados"], 1)


# --- Task 7: archivo de ejemplo ---------------------------------------------

from django.urls import reverse  # noqa: E402

from importaciones.alumnos import construir_ejemplo_alumnos  # noqa: E402
from importaciones.parsing.alumnos import ALIAS_ALUMNOS  # noqa: E402
from importaciones.parsing.comun import normalizar_texto  # noqa: E402
from tenants.models import Perfil  # noqa: E402


class EjemploAlumnosTests(BaseAlumnosTests):
    def _libro(self):
        return construir_ejemplo_alumnos()

    def test_la_vista_responde_el_xlsx(self):
        Perfil.objects.create(usuario=self.usuario, gimnasio=self.gimnasio, rol=Perfil.Rol.STAFF)
        self.client.login(username="staff", password="clave12345")
        r = self.client.get(reverse("importaciones:alumnos_ejemplo"))
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        self.assertIn("ejemplo-alumnos.xlsx", r["Content-Disposition"])
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["Alumnos", "Cómo llenarlo"])

    def test_cada_encabezado_se_detecta_como_su_campo(self):
        esperado = {
            "Nombre": "nombre", "Apellido": "apellido", "Email": "email",
            "Teléfono": "telefono", "Fecha de nacimiento": "fecha_nacimiento",
            "Sexo": "sexo", "Estado": "estado", "Inicio del ciclo": "inicio_ciclo",
            "Actividad física previa": "actividad_fisica_previa",
            "Frecuencia de actividad": "frecuencia_actividad_previa",
            "Deportes que practica": "deportes_practica",
            "Tiene discapacidad": "tiene_discapacidad",
            "Detalle discapacidad": "discapacidad_detalle",
            "Tiene enfermedad crónica": "tiene_enfermedad_cronica",
            "Detalle enfermedad crónica": "enfermedad_cronica_detalle",
            "Observaciones": "observaciones",
        }
        encabezados = [c.value for c in self._libro()["Alumnos"][1]]
        self.assertEqual(sorted(encabezados), sorted(esperado))
        buffer = io.BytesIO()
        self._libro().save(buffer)
        importacion = previsualizar_importacion_alumnos(
            gimnasio=self.gimnasio,
            archivo=SimpleUploadedFile("e.xlsx", buffer.getvalue()),
            usuario=self.usuario,
        )
        detectadas = importacion.resultado["columnas_detectadas"]
        self.assertEqual({v: k for k, v in detectadas.items()}, esperado)

    def test_los_textos_de_duplicados_dicen_las_tres_reglas(self):
        ayuda = self._libro()["Cómo llenarlo"]
        hoja = "\n".join(str(c.value) for fila in ayuda.iter_rows() for c in fila if c.value)
        Perfil.objects.create(usuario=self.usuario, gimnasio=self.gimnasio, rol=Perfil.Rol.STAFF)
        self.client.login(username="staff", password="clave12345")
        pantalla = self.client.get(reverse("importaciones:alumnos_subir")).content.decode()
        import re
        for texto in (hoja, pantalla):
            texto = " ".join(re.sub(r"<[^>]+>", "", texto).split())
            self.assertRegex(texto, r"(?i)email o teléfono y (el )?(mismo )?nombre")
            # La regla de «sin contacto» mira LOS DOS lados: la fila y el
            # alumno ya cargado (el que se cargó a mano solo con el nombre).
            self.assertIn("a la fila o al alumno ya cargado le falta el email y el teléfono", texto)
            self.assertRegex(texto, r"padres e hijos\), sí se crea")

    def test_dos_filas_de_ejemplo_una_con_inicio_y_otra_sin(self):
        ws = self._libro()["Alumnos"]
        self.assertEqual(ws.max_row, 3)
        col = [c.value for c in ws[1]].index("Inicio del ciclo") + 1
        valores = [ws.cell(row=f, column=col).value for f in (2, 3)]
        self.assertEqual(sum(1 for v in valores if not v), 1)

    def test_el_ejemplo_pasa_por_el_importador(self):
        buffer = io.BytesIO()
        self._libro().save(buffer)
        archivo = SimpleUploadedFile("ejemplo.xlsx", buffer.getvalue())
        importacion = previsualizar_importacion_alumnos(
            gimnasio=self.gimnasio, archivo=archivo, usuario=self.usuario
        )
        r = importacion.resultado
        self.assertEqual(len(r["a_crear"]), 0)
        self.assertEqual(len(r["filas_invalidas"]), 2)
        self.assertTrue(all("ejemplo" in f["motivo"].lower() for f in r["filas_invalidas"]))
        self.assertEqual(len(r["columnas_detectadas"]), 16)
        self.assertEqual(r["columnas_ignoradas"], [])

    def test_la_ayuda_dice_el_rango_y_acepta_csv(self):
        ayuda = self._libro()["Cómo llenarlo"]
        texto = "\n".join(str(c.value) for fila in ayuda.iter_rows() for c in fila if c.value)
        self.assertIn("27 días", texto)
        self.assertIn("60", texto)
        self.assertIn(".csv", texto)
        self.assertIn("dd/mm/aaaa", texto)


# --- Task 8: vistas -----------------------------------------------------------

from importaciones.forms import MAX_BYTES_ARCHIVO_ALUMNOS  # noqa: E402


class VistasAlumnosBase(BaseAlumnosTests):
    def setUp(self):
        super().setUp()
        Perfil.objects.create(usuario=self.usuario, gimnasio=self.gimnasio, rol=Perfil.Rol.STAFF)
        self.client.login(username="staff", password="clave12345")

    def _subir(self, filas, encabezados=ENCABEZADOS):
        return self.client.post(
            reverse("importaciones:alumnos_subir"),
            {"archivo": _archivo(filas, encabezados)},
        )

    def _importacion(self, filas=(("Juan", "Pérez", "", ""),)):
        return previsualizar_importacion_alumnos(
            gimnasio=self.gimnasio, archivo=_archivo(filas), usuario=self.usuario
        )


class SubirAlumnosViewTests(VistasAlumnosBase):
    def test_get_explica_el_formato_y_linkea_el_ejemplo_sin_boost(self):
        r = self.client.get(reverse("importaciones:alumnos_subir"))
        self.assertContains(r, 'href="%s" hx-boost="false"' % reverse("importaciones:alumnos_ejemplo"))
        self.assertContains(r, 'enctype="multipart/form-data" hx-boost="false"')
        self.assertContains(r, "27 días")
        self.assertContains(r, ".csv")

    def test_no_promete_sin_cobro_si_el_archivo_trae_inicio_del_ciclo(self):
        """Con la columna «Inicio del ciclo» el alumno SÍ se cobra: la pantalla
        no puede decir «sin cobro» sin condición."""
        import re

        r = self.client.get(reverse("importaciones:alumnos_subir"))
        texto = " ".join(re.sub(r"<[^>]+>", "", r.content.decode()).split())
        self.assertNotIn("sin cobro", texto)
        self.assertIn("Con ella, se le empieza a cobrar desde esa fecha.", texto)

    def test_subir_redirige_al_preview(self):
        r = self._subir([("Juan", "Pérez", "", "")])
        imp = Importacion.objects.get(tipo=Importacion.Tipo.ALUMNOS)
        self.assertRedirects(
            r, reverse("importaciones:alumnos_preview", args=[imp.pk]), fetch_redirect_response=False
        )

    def test_archivo_invalido_da_200_con_error_visible(self):
        r = self.client.post(
            reverse("importaciones:alumnos_subir"),
            {"archivo": SimpleUploadedFile("x.xlsx", b"no es un excel")},
        )
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "config-error")

    def test_extension_no_permitida(self):
        r = self.client.post(
            reverse("importaciones:alumnos_subir"),
            {"archivo": SimpleUploadedFile("x.pdf", b"%PDF")},
        )
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "config-error")

    def test_archivo_demasiado_grande_es_error_de_campo(self):
        grande = SimpleUploadedFile("x.csv", b"a" * (MAX_BYTES_ARCHIVO_ALUMNOS + 1))
        r = self.client.post(reverse("importaciones:alumnos_subir"), {"archivo": grande})
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "config-error")
        self.assertContains(r, "más de 2 MB")
        self.assertFalse(Importacion.objects.exists())

    def test_el_tope_de_tamano_es_2_mb(self):
        # 1000 filas con las 16 columnas llenas pesan ~76 KB: 2 MB sobra, y 5
        # dejaban pasar un xlsx de ~72.000 filas (22 s y ~570 MB en el worker).
        self.assertEqual(MAX_BYTES_ARCHIVO_ALUMNOS, 2 * 1024 * 1024)

    def test_demasiadas_filas_dice_mas_de_n(self):
        filas = [(f"N{i}", f"A{i}", "", "") for i in range(1001)]
        r = self._subir(filas)
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "más de 1000 filas con datos")

    def test_acepta_csv(self):
        csv_ = SimpleUploadedFile("a.csv", "Nombre,Apellido\nJuan,Pérez\n".encode("utf-8"))
        r = self.client.post(reverse("importaciones:alumnos_subir"), {"archivo": csv_})
        self.assertEqual(r.status_code, 302)

    def test_demo_bloqueada(self):
        self.gimnasio.es_demo = True
        self.gimnasio.save()
        r = self.client.get(reverse("importaciones:alumnos_subir"))
        self.assertContains(r, "En la cuenta de prueba no se importan alumnos")
        self.assertContains(r, reverse("importaciones:alumnos_ejemplo"))
        r = self._subir([("Juan", "Pérez", "", "")])
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "cuenta de prueba")
        self.assertFalse(Importacion.objects.exists())

    def test_alumno_recibe_403(self):
        imp = self._importacion()
        otro = User.objects.create_user(username="al", password="clave12345")
        Perfil.objects.create(usuario=otro, gimnasio=self.gimnasio, rol=Perfil.Rol.ALUMNO)
        self.client.login(username="al", password="clave12345")
        self.assertEqual(self.client.get(reverse("importaciones:alumnos_subir")).status_code, 403)
        self.assertEqual(
            self.client.get(reverse("importaciones:alumnos_preview", args=[imp.pk])).status_code, 403
        )


class PreviewAlumnosViewTests(VistasAlumnosBase):
    def test_muestra_fechas_dd_mm_aaaa_y_avisos(self):
        hoy = timezone.localdate()
        inicio = hoy + timedelta(days=3)
        imp = self._importacion([("Juan", "Pérez", "", "")])
        r = imp.resultado
        r["a_crear"][0]["datos"]["fecha_nacimiento"] = "1990-05-03"
        r["a_crear"][0]["datos"]["fecha_inicio_ciclo"] = inicio.isoformat()
        r["a_crear"][0]["avisos"] = ["Aviso de prueba"]
        r["avisos_lectura"] = ["Dejé de leer en la fila 9"]
        r["columnas_ignoradas"] = ["Plan"]
        Importacion.objects.filter(pk=imp.pk).update(resultado=r)
        resp = self.client.get(reverse("importaciones:alumnos_preview", args=[imp.pk]))
        self.assertContains(resp, "03/05/1990")
        self.assertContains(resp, inicio.strftime("%d/%m/%Y"))
        self.assertNotContains(resp, "1990-05-03")
        self.assertContains(resp, "Aviso de prueba")
        self.assertContains(resp, "Dejé de leer en la fila 9")
        self.assertContains(resp, "Plan")
        self.assertIn("no-store", resp["Cache-Control"])

    def test_aviso_de_nacen_vencidas(self):
        imp = self._importacion()
        r = imp.resultado
        r["nacen_vencidas"] = 3
        Importacion.objects.filter(pk=imp.pk).update(resultado=r)
        resp = self.client.get(reverse("importaciones:alumnos_preview", args=[imp.pk]))
        self.assertContains(resp, "3 alumnos van a nacer con la cuota vencida")

    def test_ya_existe_linkea_a_la_ficha_sin_boost(self):
        existente = self._alumno("Ana", "Gómez", email="a@x.com")
        imp = self._importacion([("Ana", "Gómez", "a@x.com", "")])
        resp = self.client.get(reverse("importaciones:alumnos_preview", args=[imp.pk]))
        self.assertContains(
            resp, 'href="%s" hx-boost="false"' % reverse("alumnos:detalle", args=[existente.pk])
        )

    def _boton_confirmar(self, resp):
        import re
        return re.search(r'<button[^>]*id="boton-confirmar"[^>]*>', resp.content.decode()).group(0)

    def test_sin_filas_a_crear_el_boton_viene_deshabilitado(self):
        imp = self._importacion([("Juan", "", "", "")])
        resp = self.client.get(reverse("importaciones:alumnos_preview", args=[imp.pk]))
        self.assertContains(resp, "Falta el apellido")
        self.assertIn("disabled", self._boton_confirmar(resp))

    def test_con_filas_a_crear_el_boton_no_viene_deshabilitado(self):
        imp = self._importacion([("Juan", "Pérez", "", "")])
        resp = self.client.get(reverse("importaciones:alumnos_preview", args=[imp.pk]))
        self.assertNotIn("disabled", self._boton_confirmar(resp))

    def test_el_contenedor_no_entra_al_historial_de_htmx(self):
        imp = self._importacion()
        resp = self.client.get(reverse("importaciones:alumnos_preview", args=[imp.pk]))
        self.assertContains(resp, 'class="contenido--ancho" hx-history="false"')

    def test_nacen_vencidas_concuerda_en_singular_y_plural(self):
        imp = self._importacion()
        for n, frase in (
            (1, "1 alumno va a nacer con la cuota vencida y se le pausa el acceso"),
            (2, "2 alumnos van a nacer con la cuota vencida y se les pausa el acceso"),
        ):
            r = imp.resultado
            r["nacen_vencidas"] = n
            Importacion.objects.filter(pk=imp.pk).update(resultado=r)
            resp = self.client.get(reverse("importaciones:alumnos_preview", args=[imp.pk]))
            self.assertIn(frase, " ".join(resp.content.decode().split()))

    def test_muestra_sexo_estado_contacto_y_columnas_detectadas(self):
        imp = previsualizar_importacion_alumnos(
            gimnasio=self.gimnasio,
            archivo=_archivo(
                [("Ana", "Gómez", "ana@x.com", "1155550000", "F", "Baja"),
                 ("Luis", "Díaz", "", "", "M", "")],
                ("Nombre", "Apellido", "Email", "Celular", "Sexo", "Estado"),
            ),
            usuario=self.usuario,
        )
        resp = self.client.get(reverse("importaciones:alumnos_preview", args=[imp.pk]))
        self.assertContains(resp, "Femenino")
        self.assertContains(resp, "Masculino")
        self.assertContains(resp, "Inactivo")
        self.assertContains(resp, "ana@x.com")
        self.assertContains(resp, "1155550000")
        self.assertContains(resp, "Teléfono ← «Celular»")
        self.assertContains(resp, "Se leyeron 2 filas")
        self.assertNotContains(resp, "telefono ←")

    def test_preview_de_otro_gimnasio_404(self):
        otro = Gimnasio.objects.create(nombre="Otro", slug="otro")
        imp = previsualizar_importacion_alumnos(
            gimnasio=otro, archivo=_archivo([("Juan", "Pérez", "", "")]), usuario=self.usuario
        )
        self.assertEqual(
            self.client.get(reverse("importaciones:alumnos_preview", args=[imp.pk])).status_code, 404
        )

    def test_preview_de_otro_tipo_404(self):
        imp = self._importacion()
        Importacion.objects.filter(pk=imp.pk).update(tipo=Importacion.Tipo.BIBLIOTECA)
        self.assertEqual(
            self.client.get(reverse("importaciones:alumnos_preview", args=[imp.pk])).status_code, 404
        )

    def test_post_confirma_y_redirige_al_listado(self):
        imp = self._importacion([("Juan", "Pérez", "", "")])
        with self.captureOnCommitCallbacks(execute=True):
            r = self.client.post(reverse("importaciones:alumnos_preview", args=[imp.pk]))
        self.assertEqual(Alumno.objects.count(), 1)
        self.assertRedirects(
            r, reverse("alumnos:listado") + "?sin_fecha=1", fetch_redirect_response=False
        )
        textos = [str(m) for m in r.wsgi_request._messages]
        self.assertTrue(any("1 alumno" in t for t in textos))
        self.assertTrue(any("sin fecha de inicio" in t and "Fijar fecha de inicio" in t for t in textos))

    def test_post_sin_alumnos_sin_fecha_va_al_listado_pelado(self):
        inicio = timezone.localdate().strftime("%d/%m/%Y")
        imp = previsualizar_importacion_alumnos(
            gimnasio=self.gimnasio,
            archivo=_archivo([("Juan", "Pérez", inicio)], ("Nombre", "Apellido", "Inicio del ciclo")),
            usuario=self.usuario,
        )
        with self.captureOnCommitCallbacks(execute=True):
            r = self.client.post(reverse("importaciones:alumnos_preview", args=[imp.pk]))
        self.assertRedirects(r, reverse("alumnos:listado"), fetch_redirect_response=False)

    def test_post_sobre_ya_confirmada_redirige_con_mensaje(self):
        imp = self._importacion()
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(reverse("importaciones:alumnos_preview", args=[imp.pk]))
        for metodo in (self.client.post, self.client.get):
            r = metodo(reverse("importaciones:alumnos_preview", args=[imp.pk]))
            self.assertRedirects(r, reverse("alumnos:listado"), fetch_redirect_response=False)
            self.assertIn("Esta importación ya se procesó.", [str(m) for m in r.wsgi_request._messages])
        self.assertEqual(Alumno.objects.count(), 1)

    def test_importacion_invalida_en_el_post_da_200_con_mensaje(self):
        imp = self._importacion()
        with mock.patch(
            "importaciones.views.confirmar_importacion_alumnos",
            side_effect=ImportacionInvalida("Algo falló"),
        ):
            r = self.client.post(reverse("importaciones:alumnos_preview", args=[imp.pk]))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Algo falló")


class DescartarAlumnosTests(VistasAlumnosBase):
    def test_descartar_redirige_a_subir_y_achica_el_resultado(self):
        imp = self._importacion()
        archivo = imp.archivo.name
        with self.captureOnCommitCallbacks(execute=True):
            r = self.client.post(reverse("importaciones:alumnos_descartar", args=[imp.pk]))
        self.assertRedirects(r, reverse("importaciones:alumnos_subir"), fetch_redirect_response=False)
        imp.refresh_from_db()
        self.assertEqual(imp.estado, Importacion.Estado.DESCARTADA)
        self.assertEqual(set(imp.resultado), {"version", "resumen"})
        self.assertEqual(imp.resultado["resumen"]["creados"], 0)
        self.assertFalse(imp.archivo)
        self.assertFalse(default_storage.exists(archivo))

    def test_descartar_otra_importacion_404(self):
        imp = self._importacion()
        otro = Gimnasio.objects.create(nombre="Otro", slug="otro")
        Importacion.objects.filter(pk=imp.pk).update(gimnasio=otro)
        r = self.client.post(reverse("importaciones:alumnos_descartar", args=[imp.pk]))
        self.assertEqual(r.status_code, 404)


class BotonEnElListadoTests(VistasAlumnosBase):
    def test_el_listado_muestra_el_boton(self):
        r = self.client.get(reverse("alumnos:listado"))
        self.assertContains(r, reverse("importaciones:alumnos_subir"))
        self.assertContains(r, "Importar desde Excel")


class DescartarDobleClickTests(VistasAlumnosBase):
    def test_doble_post_de_alumnos_redirige_con_mensaje(self):
        imp = self._importacion()
        url = reverse("importaciones:alumnos_descartar", args=[imp.pk])
        self.client.post(url)
        r = self.client.post(url)
        self.assertRedirects(r, reverse("importaciones:alumnos_subir"), fetch_redirect_response=False)
        self.assertIn("Esta importación ya se procesó.", [str(m) for m in r.wsgi_request._messages])

    def test_plantillas_y_biblioteca_siguen_dando_404_si_ya_no_estan_en_revision(self):
        for tipo, nombre in (
            (Importacion.Tipo.BIBLIOTECA, "biblioteca_descartar"),
            (Importacion.Tipo.PLANTILLAS, "plantillas_descartar"),
        ):
            imp = Importacion.objects.create(
                gimnasio=self.gimnasio, tipo=tipo, resultado={},
                estado=Importacion.Estado.DESCARTADA, creado_por=self.usuario,
            )
            r = self.client.post(reverse(f"importaciones:{nombre}", args=[imp.pk]))
            self.assertEqual(r.status_code, 404, tipo)

    def test_el_boton_de_descartar_se_deshabilita_al_enviar(self):
        imp = self._importacion()
        resp = self.client.get(reverse("importaciones:alumnos_preview", args=[imp.pk]))
        self.assertContains(resp, 'id="boton-descartar"')
        self.assertContains(resp, 'getElementById("importacion-descartar").addEventListener("submit"')
